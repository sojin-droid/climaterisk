"""Asset import from a spreadsheet template — three granularities, one parser.

The EIOPA CLIMADA-App ships exposure templates at three levels (lat/lon, NUTS2, country).
The equivalents here are ``latlon`` (a located asset per row), ``sigungu`` (a 시군구 per
row, placed at its official representative point) and ``sido`` (a 시도 per row, likewise).
A workbook or CSV in any of them becomes :class:`~climaterisk.core.entities.Asset` rows
for the session portfolio; nothing is priced here.

Rules, all deliberate:

* **Nothing is invented.** A municipality row without a value gets ``value = 0`` and is
  therefore hazard screening only downstream (``NO_EXPOSURE_DATA``); a lat/lon row
  without a value is refused, because a located facility with no value is a typo, not a
  screening request.
* **Ambiguity is an error, not a guess.** 시군구 names repeat across provinces (중구 ×6,
  동구 ×6, …): such a row needs ``municipality_id`` or ``province_name``.
* **Errors are line-numbered** and the whole file is refused when any row is wrong, so a
  partially imported portfolio never exists.
* **Re-uploads are idempotent** under ``mode="merge"``: a row whose id already exists
  replaces that asset; ``mode="replace"`` swaps the whole portfolio for the file.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook

from climaterisk.core.entities import Asset
from climaterisk.core.enums import Sector

TEMPLATE_KINDS: tuple[str, ...] = ("latlon", "sigungu", "sido")
IMPORT_MODES: tuple[str, ...] = ("merge", "replace")

#: Canonical field -> accepted header spellings (case-insensitive, trimmed).
_ALIASES: dict[str, tuple[str, ...]] = {
    "facility_id": ("facility_id", "id", "asset_id"),
    "facility_name": ("facility_name", "name", "asset_name"),
    "latitude": ("latitude", "lat"),
    "longitude": ("longitude", "lon", "lng"),
    "asset_value_usd": ("asset_value_usd", "value", "asset_value", "value_usd"),
    "asset_value_currency": ("asset_value_currency", "currency"),
    "property_type": ("property_type", "type", "sector"),
    "headcount": ("headcount", "people", "employees"),
    "municipality_id": ("municipality_id", "sgis_id", "admin_id"),
    "municipality_name": ("municipality_name", "municipality", "sigungu", "시군구", "sido", "시도"),
    "province_name": ("province_name", "province", "시도명"),
}

#: Per template: the columns in order, which are required, and an example row.
TEMPLATES: dict[str, dict[str, Any]] = {
    "latlon": {
        "title": "Assets by coordinates",
        "columns": (
            "facility_id",
            "facility_name",
            "latitude",
            "longitude",
            "asset_value_usd",
            "asset_value_currency",
            "property_type",
            "headcount",
        ),
        "required": ("facility_id", "latitude", "longitude", "asset_value_usd"),
        "examples": (
            ("KR-001", "Seoul office", 37.5665, 126.978, 100000000, "USD", "office", 120),
            ("KR-002", "Busan warehouse", 35.1796, 129.0756, 60000000, "USD", "warehouse", 35),
        ),
        "notes": (
            "One located asset per row. latitude/longitude in decimal degrees (WGS84).",
            "asset_value_usd is required: a located asset without a value is refused.",
            "property_type: a sector id (e.g. real_estate, steel) or a JRC-mappable label "
            "(office, warehouse, residential, …); unknown labels keep the asset's default "
            "sector and are carried as property_type.",
            "headcount is optional (people on site; exposure for heat mortality).",
        ),
    },
    "sigungu": {
        "title": "Municipalities (시군구) by official representative point",
        "columns": (
            "municipality_id",
            "municipality_name",
            "province_name",
            "asset_value_usd",
            "asset_value_currency",
            "property_type",
            "facility_id",
        ),
        "required": (),
        "examples": (
            ("SGIS-32010", "춘천시", "강원특별자치도", 50000000, "USD", "office", ""),
            ("", "중구", "서울특별시", "", "", "", ""),
        ),
        "notes": (
            "One 시군구 per row, placed at its official representative point (Statistics "
            "Korea SGIS 2025 2Q boundary interior point — not the city hall, not an area).",
            "Give municipality_id (SGIS-…) or municipality_name; a name that repeats across "
            "provinces (중구, 동구, 남구, 서구, 북구, 강서구, 고성군) also needs province_name.",
            "asset_value_usd is optional: without it the row is hazard screening only — "
            "no value is ever estimated for a municipality.",
            "facility_id is optional; the municipality_id is used when it is blank, so a "
            "re-upload of the same file replaces the same rows.",
        ),
    },
    "sido": {
        "title": "Provinces (시도) by official representative point",
        "columns": (
            "municipality_id",
            "municipality_name",
            "asset_value_usd",
            "asset_value_currency",
            "property_type",
            "facility_id",
        ),
        "required": (),
        "examples": (
            ("SGIS-11", "서울특별시", 250000000, "USD", "office", ""),
            ("", "부산광역시", "", "", "", ""),
        ),
        "notes": (
            "One 시도 per row, placed at the 시도 unit's own representative point (an "
            "interior point of its boundary). It is a screening anchor for the whole "
            "province, not an aggregate over its municipalities.",
            "Give municipality_id (SGIS-…) or the official 시도 name.",
            "asset_value_usd is optional (screening only without it).",
        ),
    },
}

_SECTOR_IDS = {s.value for s in Sector}


@dataclass
class ImportReport:
    """What an import did, line by line where it matters."""

    kind: str
    mode: str
    rows_read: int = 0
    added: list[str] = field(default_factory=list)
    replaced: list[str] = field(default_factory=list)
    removed: list[str] = field(default_factory=list)  # mode=replace: ids no longer present
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "kind": self.kind,
            "mode": self.mode,
            "rows_read": self.rows_read,
            "added": list(self.added),
            "replaced": list(self.replaced),
            "removed": list(self.removed),
            "warnings": list(self.warnings),
            "errors": list(self.errors),
        }


class AssetImportError(ValueError):
    """The file was refused; ``errors`` are line-numbered and complete."""

    def __init__(self, errors: list[str]) -> None:
        super().__init__("; ".join(errors[:5]) + (" …" if len(errors) > 5 else ""))
        self.errors = errors


# --------------------------------------------------------------------------- #
# Templates                                                                    #
# --------------------------------------------------------------------------- #
def template_workbook(kind: str) -> bytes:
    """The downloadable xlsx for ``kind``: a README sheet and an ``assets`` sheet."""
    if kind not in TEMPLATES:
        raise KeyError(f"unknown template kind {kind!r}; one of {TEMPLATE_KINDS}")
    spec = TEMPLATES[kind]
    wb = Workbook()
    readme = wb.active
    assert readme is not None
    readme.title = "README"
    readme.append([f"climaterisk asset template — {spec['title']}"])
    readme.append(["Fill the 'assets' sheet; keep the header row. Example rows may be deleted."])
    readme.append([])
    readme.append(["column", "required", "accepted header spellings"])
    for col in spec["columns"]:
        readme.append([col, "yes" if col in spec["required"] else "no", ", ".join(_ALIASES[col])])
    readme.append([])
    for note in spec["notes"]:
        readme.append([note])
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["C"].width = 60
    ws = wb.create_sheet("assets")
    ws.append(list(spec["columns"]))
    for row in spec["examples"]:
        ws.append([None if v == "" else v for v in row])
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Reading                                                                      #
# --------------------------------------------------------------------------- #
def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_table(filename: str, content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    """``(header, rows)`` from an xlsx (sheet ``assets`` or the first sheet) or a CSV.

    Cells are returned as trimmed strings; a blank row is skipped. The format is taken from
    the file name's extension, then from the content (xlsx files start with ``PK``).
    """
    name = (filename or "").lower()
    is_xlsx = name.endswith((".xlsx", ".xlsm")) or (
        not name.endswith(".csv") and content[:2] == b"PK"
    )
    if is_xlsx:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb["assets"] if "assets" in wb.sheetnames else wb[wb.sheetnames[0]]
        lines = [[_cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = content.decode("utf-8-sig")
        lines = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]
    lines = [r for r in lines if any(r)]
    if not lines:
        return [], []
    header = [h.strip() for h in lines[0]]
    rows = [
        {header[i]: (r[i] if i < len(r) else "") for i in range(len(header))} for r in lines[1:]
    ]
    return header, rows


def _pick_columns(header: list[str]) -> dict[str, str]:
    """Canonical field -> actual header for every alias found (first match wins)."""
    lower = {h.lower(): h for h in header}
    picked: dict[str, str] = {}
    for fieldname, names in _ALIASES.items():
        for n in names:
            if n.lower() in lower:
                picked[fieldname] = lower[n.lower()]
                break
    return picked


def _sector_and_property(label: str) -> tuple[Sector, str | None]:
    """A sector id becomes the sector; any other label is carried as ``property_type``."""
    if not label:
        return Sector.REAL_ESTATE, None
    key = label.strip().lower()
    if key in _SECTOR_IDS:
        return Sector(key), None
    return Sector.REAL_ESTATE, label.strip()


def _number(raw: str, what: str, line: int, errors: list[str]) -> float | None:
    if raw == "":
        return None
    try:
        return float(raw.replace(",", ""))
    except ValueError:
        errors.append(f"line {line}: {what} is not a number ({raw!r})")
        return None


def parse_assets(
    filename: str, content: bytes, kind: str, existing_ids: set[str] | None = None
) -> tuple[list[Asset], ImportReport]:
    """Assets from a template file. Raises :class:`AssetImportError` when any row is wrong."""
    if kind not in TEMPLATES:
        raise AssetImportError([f"unknown template kind {kind!r}; one of {TEMPLATE_KINDS}"])
    header, rows = read_table(filename, content)
    report = ImportReport(kind=kind, mode="")
    if not header:
        raise AssetImportError(["the file has no header row"])
    picked = _pick_columns(header)
    spec = TEMPLATES[kind]
    missing = [c for c in spec["required"] if c not in picked]
    if missing:
        raise AssetImportError([f"missing required column(s) {missing}; header was {header}"])
    if kind != "latlon" and "municipality_id" not in picked and "municipality_name" not in picked:
        raise AssetImportError(
            [f"a {kind} template needs municipality_id or municipality_name; header was {header}"]
        )

    errors: list[str] = []
    assets: list[Asset] = []
    seen: set[str] = set()
    for i, raw in enumerate(rows, start=2):
        row = {k: raw.get(col, "") for k, col in picked.items()}
        asset = (
            _latlon_asset(row, i, errors)
            if kind == "latlon"
            else _admin_asset(row, i, kind, errors)
        )
        if asset is None:
            continue
        if asset.id in seen:
            errors.append(f"line {i}: duplicate id {asset.id!r} within the file")
            continue
        seen.add(asset.id)
        assets.append(asset)
    report.rows_read = len(rows)
    if errors:
        raise AssetImportError(errors)
    if not assets:
        raise AssetImportError(["the file has a header but no data rows"])
    for a in assets:
        (report.replaced if existing_ids and a.id in existing_ids else report.added).append(a.id)
    return assets, report


def _latlon_asset(row: dict[str, str], line: int, errors: list[str]) -> Asset | None:
    fid = row.get("facility_id", "")
    if not fid:
        errors.append(f"line {line}: empty facility_id")
        return None
    lat = _number(row.get("latitude", ""), "latitude", line, errors)
    lon = _number(row.get("longitude", ""), "longitude", line, errors)
    value = _number(row.get("asset_value_usd", ""), "asset_value_usd", line, errors)
    if lat is None or lon is None:
        if row.get("latitude", "") == "" or row.get("longitude", "") == "":
            errors.append(f"line {line}: latitude and longitude are required")
        return None
    if not (-90.0 <= lat <= 90.0 and -180.0 <= lon <= 180.0):
        errors.append(f"line {line}: latitude/longitude out of range ({lat}, {lon})")
        return None
    if value is None:
        if row.get("asset_value_usd", "") == "":
            errors.append(f"line {line}: asset_value_usd is required for a located asset")
        return None
    if value < 0:
        errors.append(f"line {line}: asset_value_usd is negative ({value})")
        return None
    headcount = _number(row.get("headcount", ""), "headcount", line, errors)
    if headcount is not None and headcount < 0:
        errors.append(f"line {line}: headcount is negative ({headcount})")
        return None
    sector, ptype = _sector_and_property(row.get("property_type", ""))
    props: dict[str, float | str | bool] = {"import_template": "latlon"}
    if ptype:
        props["property_type"] = ptype
    return Asset(
        id=fid,
        name=row.get("facility_name") or fid,
        lat=lat,
        lon=lon,
        sector=sector,
        value=value,
        currency=row.get("asset_value_currency") or "USD",
        headcount=int(headcount) if headcount is not None else None,
        properties=props,
    )


def _admin_asset(row: dict[str, str], line: int, kind: str, errors: list[str]) -> Asset | None:
    from climaterisk.physical_risk.municipalities import load_municipalities

    tier = "SIDO" if kind == "sido" else "SIGUNGU"
    units = [m for m in load_municipalities() if m.tier == tier]
    mid, name, prov = (
        row.get("municipality_id", ""),
        row.get("municipality_name", ""),
        row.get("province_name", ""),
    )
    if mid:
        hits = [m for m in units if m.municipality_id == mid]
        if not hits:
            errors.append(f"line {line}: unknown {tier.lower()} municipality_id {mid!r}")
            return None
    elif name:
        hits = [m for m in units if m.municipality_name == name]
        if prov:
            hits = [m for m in hits if m.province_name == prov]
        if not hits:
            errors.append(
                f"line {line}: no {tier.lower()} named {name!r}" + (f" in {prov!r}" if prov else "")
            )
            return None
        if len(hits) > 1:
            provinces = sorted(m.province_name for m in hits)
            errors.append(
                f"line {line}: {name!r} is ambiguous ({len(hits)} units: {provinces}); "
                "give municipality_id or province_name"
            )
            return None
    else:
        errors.append(f"line {line}: municipality_id or municipality_name is required")
        return None
    m = hits[0]
    value = _number(row.get("asset_value_usd", ""), "asset_value_usd", line, errors)
    if value is not None and value < 0:
        errors.append(f"line {line}: asset_value_usd is negative ({value})")
        return None
    sector, ptype = _sector_and_property(row.get("property_type", ""))
    props: dict[str, float | str | bool] = {
        "import_template": kind,
        "municipality_id": m.municipality_id,
        "municipality_name": m.municipality_name,
        "municipality_level": m.municipality_level,
        "province_name": m.province_name,
        "point_type": m.point_type,
    }
    if ptype:
        props["property_type"] = ptype
    return Asset(
        id=row.get("facility_id") or m.municipality_id,
        name=m.municipality_name,
        lat=m.latitude,
        lon=m.longitude,
        sector=sector,
        value=value if value is not None else 0.0,
        currency=row.get("asset_value_currency") or "USD",
        properties=props,
    )


# --------------------------------------------------------------------------- #
# Applying to a portfolio                                                      #
# --------------------------------------------------------------------------- #
def merge_assets(
    existing: list[Asset], incoming: list[Asset], mode: str
) -> tuple[list[Asset], list[str]]:
    """``(assets, removed_ids)`` after applying ``incoming`` under ``mode``.

    ``merge``: an incoming id replaces the existing asset in place, new ids are appended,
    everything else is kept. ``replace``: the result is exactly ``incoming``.
    """
    if mode not in IMPORT_MODES:
        raise ValueError(f"mode must be one of {IMPORT_MODES}, got {mode!r}")
    if mode == "replace":
        new_ids = {a.id for a in incoming}
        return list(incoming), [a.id for a in existing if a.id not in new_ids]
    by_id = {a.id: a for a in incoming}
    out = [by_id.pop(a.id, a) for a in existing]
    out += [a for a in incoming if a.id in by_id]
    return out, []
