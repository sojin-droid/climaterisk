"""V0.2 — municipality representative points, five-hazard coverage, spatial resolution, exports.

CLIMADA-free. The rules under test:

* the bundled dataset is valid (unique ids, coordinates inside Korea, known levels and
  point types, provenance stated) and its representative points are boundary interior
  points — never city halls, never area aggregates;
* a municipality carries no asset value unless the caller supplies one; a screening row
  is ``NO_EXPOSURE_DATA`` with hazard intensity and **null** money — never ``0``;
* the five-hazard coverage table and the resolution registry are *generated* from the
  readiness table and the extended-hazard evidence, not typed in; drought and sea-level
  rise are ``NOT_READY`` with a stated reason;
* the municipality Excel has the seven required sheets in order, an empty cell for an
  unpriced field, and the municipality filename; the facility workbook is unchanged;
* the API validates ids and values, serves the dataset, and the request builder never
  invents a value.
"""

from __future__ import annotations

import csv
import io
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
if str(REPO / "src") not in sys.path:
    sys.path.insert(0, str(REPO / "src"))

from climaterisk.engines.base import PhysicalRiskModelsRequest  # noqa: E402
from climaterisk.physical_risk import coverage as cov  # noqa: E402
from climaterisk.physical_risk import display_copy as dc  # noqa: E402
from climaterisk.physical_risk import excel_export as xl  # noqa: E402
from climaterisk.physical_risk import municipalities as mu  # noqa: E402
from climaterisk.physical_risk import results_table as rt  # noqa: E402
from climaterisk.physical_risk.metrics import CalculationStatus, ModelId, ResultRow  # noqa: E402
from climaterisk.physical_risk.models import HAZARD_KEYS, READINESS  # noqa: E402

G, C, K = ModelId.GLOBAL_BASELINE.value, ModelId.DATA_API_COUNTRY.value, ModelId.KOREA_LOCAL.value
GWANGJU, SEOUL, BUSAN, CHUNCHEON, INCHEON = "SGIS-24", "SGIS-11", "SGIS-21", "SGIS-32010", "SGIS-23"


# --------------------------------------------------------------------------- #
# Dataset                                                                       #
# --------------------------------------------------------------------------- #
def test_dataset_loads_and_is_valid() -> None:
    ms = mu.load_municipalities()
    assert len(ms) >= 17 + 200  # 17 시도 + the SGIS 시군구 units
    assert not mu.validate([m.to_dict() for m in ms])
    assert {m.municipality_id for m in ms} >= {GWANGJU, SEOUL, BUSAN, CHUNCHEON, INCHEON}
    assert all(33.0 <= m.latitude <= 39.0 and 124.5 <= m.longitude <= 131.0 for m in ms)
    assert all(m.point_type == mu.POINT_TYPE_BOUNDARY_INTERIOR for m in ms)
    assert all(m.office_name is None and m.address is None for m in ms)  # not city halls
    assert all(m.source and m.source_version and m.source_licence for m in ms)
    assert all(m.municipality_level in mu.LEVELS for m in ms)


def test_dataset_levels_and_tiers() -> None:
    ms = {m.municipality_id: m for m in mu.load_municipalities()}
    assert ms[GWANGJU].municipality_level == "METROPOLITAN" and ms[GWANGJU].tier == "SIDO"
    assert (
        ms["SGIS-31"].municipality_name == "경기도"
        and ms["SGIS-31"].municipality_level == "PROVINCE"
    )
    assert ms[CHUNCHEON].municipality_level == "CITY" and ms[CHUNCHEON].tier == "SIGUNGU"
    assert ms[CHUNCHEON].province_name == "강원특별자치도"
    summary = mu.dataset_summary()
    assert summary["total"] == len(ms) and sum(summary["by_level"].values()) == len(ms)
    assert "interior point" in summary["representative_point_definition"]
    assert summary["warning"] == mu.REPRESENTATIVE_POINT_WARNING


def test_validate_catches_duplicates_bad_coordinates_and_unknown_levels() -> None:
    good = mu.load_municipalities()[0].to_dict()
    dup = [good, dict(good)]
    assert any("duplicate" in p for p in mu.validate(dup))
    off = dict(good, latitude=51.5, longitude=-0.1)
    assert any("outside Korea" in p for p in mu.validate([off]))
    bad_level = dict(good, municipality_level="PLANET")
    assert any("unknown level" in p for p in mu.validate([bad_level]))
    bad_point = dict(good, point_type="GUESSED")
    assert any("point_type" in p for p in mu.validate([bad_point]))
    nan = dict(good, latitude="abc")
    assert any("non-numeric" in p for p in mu.validate([nan]))


def test_read_dataset_rejects_an_invalid_file(tmp_path: Path) -> None:
    p = tmp_path / "bad.csv"
    with p.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(mu.COLUMNS))
        w.writeheader()
        row = mu.load_municipalities()[0].to_dict()
        row.pop("tier")
        w.writerow(row)
        w.writerow(row)  # duplicate id
    with pytest.raises(ValueError, match="duplicate"):
        mu.read_dataset(p)


def test_by_id_preserves_order_and_rejects_unknown() -> None:
    picked = mu.by_id([BUSAN, GWANGJU])
    assert [m.municipality_id for m in picked] == [BUSAN, GWANGJU]
    with pytest.raises(KeyError):
        mu.by_id(["SGIS-00000"])


def test_as_facility_never_invents_a_value() -> None:
    m = mu.by_id([GWANGJU])[0]
    f = m.as_facility()
    assert f["asset_value_usd"] is None and f["assessment_target"] == "MUNICIPALITY"
    assert f["facility_id"] == GWANGJU and f["municipality_name"] == "광주광역시"
    assert f["point_type"] == mu.POINT_TYPE_BOUNDARY_INTERIOR
    g = m.as_facility(80_000_000.0, "office")
    assert g["asset_value_usd"] == 80_000_000.0 and g["property_type"] == "office"


def test_representative_point_warning_is_verbatim() -> None:
    assert mu.REPRESENTATIVE_POINT_WARNING == (
        "This result represents hazard conditions at the selected municipality's representative "
        "point. It is not a municipality-wide spatial aggregation."
    )
    bundle = dc.display_bundle()["municipality"]
    assert bundle["warning"] == mu.REPRESENTATIVE_POINT_WARNING
    assert bundle["dataset"]["total"] == len(mu.load_municipalities())


# --------------------------------------------------------------------------- #
# Request builder                                                               #
# --------------------------------------------------------------------------- #
def test_request_from_municipalities_supplies_values_only_where_given() -> None:
    req = PhysicalRiskModelsRequest.from_municipalities(
        "s1", "rcp85", [GWANGJU, SEOUL], asset_values={GWANGJU: 80e6}, target_year=2040
    )
    assert req.assessment_target == "MUNICIPALITY" and req.country == "KOR"
    by = {f.facility_id: f for f in req.facilities}
    assert by[GWANGJU].asset_value_usd == 80e6 and by[SEOUL].asset_value_usd is None
    assert (
        by[SEOUL].municipality_name == "서울특별시"
        and by[SEOUL].assessment_target == "MUNICIPALITY"
    )
    payload = req.model_dump()
    assert payload["facilities"][1]["point_type"] == mu.POINT_TYPE_BOUNDARY_INTERIOR


# --------------------------------------------------------------------------- #
# Coverage + resolution registries                                              #
# --------------------------------------------------------------------------- #
def test_coverage_table_is_generated_from_the_registries() -> None:
    table = cov.coverage_table()
    assert (
        [r["hazard_key"] for r in table]
        == list(cov.COVERAGE_KEYS)
        == ["RF", "TC", "HEAT", "DROUGHT", "SLR"]
    )
    by = {r["hazard_key"]: r for r in table}
    assert by["RF"]["Status"] == "READY" and by["RF"]["Financial Loss"] == "Available"
    assert by["TC"]["Status"] == "READY"
    assert by["HEAT"]["Status"] == "HAZARD_ONLY" and by["HEAT"]["Financial Loss"] == "Not available"
    assert by["DROUGHT"]["Status"] == "NOT_READY" and by["SLR"]["Status"] == "NOT_READY"
    assert by["DROUGHT"]["Why"].startswith("A legacy drought result exists")
    assert by["SLR"]["Why"].startswith("The current implementation is configuration-based")
    assert by["HEAT"]["Why"].startswith(
        "Hazard data is available, but no applicable CLIMADA Impact Function"
    )
    # the statuses of the core hazards are the readiness table's, not a copy
    assert by["RF"]["models"] == {m: c["status"] for m, c in READINESS["RF"].items()}
    # resolution text comes from the registry values
    assert "150 arcsec" in by["RF"]["Resolution"] and "km" in by["RF"]["Resolution"]
    assert "0.05 degree" in by["HEAT"]["Resolution"]
    assert by["SLR"]["Resolution"].startswith("N/A")


def test_coverage_follows_a_changed_readiness_table() -> None:
    altered = {k: {m: dict(c) for m, c in v.items()} for k, v in READINESS.items()}
    altered["RF"][C]["status"] = "NOT_IMPLEMENTED"
    altered["RF"][G]["status"] = "NOT_IMPLEMENTED"
    by = {r["hazard_key"]: r for r in cov.coverage_table(altered)}
    assert by["RF"]["Status"] == "NOT_READY"  # no runnable cell any more


def test_coverage_lines_match_the_table() -> None:
    lines = cov.coverage_summary_lines()
    assert lines == [
        "✓ Flood — financial risk",
        "✓ Tropical Cyclone — financial risk",
        "✓ Heatwave — hazard only",
        "○ Drought — not ready",
        "○ Sea-level rise — not ready",
    ]


def test_resolution_registry_values_and_units() -> None:
    assert (
        cov.resolution_for("RF", C)["value"] == 150
        and cov.resolution_for("RF", C)["unit"] == "arcsec"
    )
    assert cov.resolution_for("TC", G)["unit_type"] == "grid cell"
    heat = cov.resolution_for("HEAT", K)
    assert heat["value"] == 0.05 and heat["unit"] == "degree" and heat["unit_type"] == "grid"
    assert cov.resolution_for("HEAT", C) is None and cov.resolution_for("DROUGHT", C) is None
    assert cov.km_class(150, "arcsec") == "~4–5 km"
    assert cov.km_class(0.05, "degree") == "~4–6 km"
    assert cov.approx_km(1.0, "index") is None
    reg = cov.resolution_registry()
    assert reg["note"] == cov.RESOLUTION_NOTE and reg["extended"]["SLR"]["status"] == "NOT_READY"
    assert "not by itself mean higher model accuracy" in cov.RESOLUTION_NOTE


def test_extended_hazards_carry_import_level_evidence() -> None:
    for key in ("DROUGHT", "SLR"):
        ext = cov.EXTENDED_HAZARDS[key]
        assert ext["status"] == "NOT_READY" and ext["financial_loss"] == "Not available"
        assert "checked" in ext["evidence"] and ext["candidate"]
    assert "ImpfDrought.from_default()" in cov.EXTENDED_HAZARDS["DROUGHT"]["evidence"]
    assert "sea_level_rise_m" in cov.EXTENDED_HAZARDS["SLR"]["evidence"]["repository"]


def test_core_readiness_table_is_untouched_by_v02() -> None:
    assert HAZARD_KEYS == ("RF", "TC", "HEAT")
    assert {k: {m: c["status"] for m, c in v.items()} for k, v in READINESS.items()} == {
        "RF": {G: "READY", C: "READY", K: "NOT_IMPLEMENTED"},
        "TC": {G: "READY", C: "READY", K: "NOT_IMPLEMENTED"},
        "HEAT": {G: "NO_HAZARD_DATA", C: "NO_HAZARD_DATA", K: "HAZARD_ONLY"},
    }
    assert dc.recommended_models() == {"RF": C, "TC": C, "HEAT": K}


def test_display_bundle_exposes_coverage_resolution_and_why() -> None:
    b = dc.display_bundle()
    assert (
        b["hazard_label"]["DROUGHT"] == "Drought" and b["hazard_label"]["SLR"] == "Sea-level rise"
    )
    assert set(b["why_unavailable"]) == {"HEAT", "DROUGHT", "SLR"}
    assert b["resolution_note"] == cov.RESOLUTION_NOTE
    assert b["municipality"]["status_short"]["NO_EXPOSURE_DATA"] == "No asset value"
    assert b["coverage_table"][0]["Hazard"] == "Flood"


# --------------------------------------------------------------------------- #
# Rows, frames, Excel                                                           #
# --------------------------------------------------------------------------- #
def _muni_row(mid: str, haz: str, model: str, **kw) -> ResultRow:  # type: ignore[no-untyped-def]
    m = mu.by_id([mid])[0]
    base = {
        "facility_id": mid,
        "facility_name": m.municipality_name,
        "hazard_type": haz,
        "model_id": model,
        "assessment_target": "MUNICIPALITY",
        "municipality_id": mid,
        "municipality_name": m.municipality_name,
        "municipality_level": m.municipality_level,
        "province_name": m.province_name,
        "point_type": m.point_type,
        "latitude": m.latitude,
        "longitude": m.longitude,
        "requested_scenario": "rcp85",
        "served_scenario": "rcp85",
        "hazard_dataset": f"synthetic_{haz}_{model}",
        "spatial_resolution": 0.05 if haz == "HW" else 150.0,
        "spatial_resolution_unit": "degree" if haz == "HW" else "arcsec",
        "spatial_unit_type": "grid" if haz == "HW" else "grid cell",
        "spatial_resolution_description": "synthetic",
    }
    base.update(kw)
    return ResultRow(**base).finalise()  # type: ignore[arg-type]


def _run_rows() -> list[ResultRow]:
    return [
        # Gwangju: valued, positive flood → High; TC low
        _muni_row(
            GWANGJU,
            "RF",
            C,
            asset_value_usd=80e6,
            hazard_intensity=8.31,
            hazard_intensity_unit="m",
            eal_usd=597_066.6,
            potential_loss_usd=5.0e6,
            return_period_years=100.0,
            impact_function_id=22,
            calculation_status="FULL",
        ),
        _muni_row(
            GWANGJU,
            "TC",
            C,
            asset_value_usd=80e6,
            hazard_intensity=55.2,
            hazard_intensity_unit="m/s",
            eal_usd=11_200.2,
            potential_loss_usd=2.0e5,
            return_period_years=100.0,
            impact_function_id=9,
            calculation_status="FULL",
        ),
        _muni_row(
            GWANGJU,
            "HW",
            K,
            asset_value_usd=80e6,
            hazard_intensity=37.8,
            hazard_intensity_unit="degC",
            calculation_status="HAZARD_ONLY",
        ),
        # Busan: valued, flood computed zero → Low + 0
        _muni_row(
            BUSAN,
            "RF",
            C,
            asset_value_usd=60e6,
            hazard_intensity=0.0,
            hazard_intensity_unit="m",
            eal_usd=0.0,
            potential_loss_usd=0.0,
            return_period_years=100.0,
            impact_function_id=21,
            calculation_status="FULL",
        ),
        _muni_row(
            BUSAN,
            "TC",
            C,
            asset_value_usd=60e6,
            hazard_intensity=60.3,
            hazard_intensity_unit="m/s",
            eal_usd=14_749.3,
            potential_loss_usd=3.0e5,
            return_period_years=100.0,
            impact_function_id=9,
            calculation_status="FULL",
        ),
        _muni_row(
            BUSAN,
            "HW",
            K,
            asset_value_usd=60e6,
            hazard_intensity=34.1,
            hazard_intensity_unit="degC",
            calculation_status="HAZARD_ONLY",
        ),
        # Seoul: no value → screening only
        _muni_row(
            SEOUL,
            "RF",
            C,
            hazard_intensity=6.97,
            hazard_intensity_unit="m",
            impact_function_id=22,
            calculation_status="NO_EXPOSURE_DATA",
            status_detail="no asset value supplied",
        ),
        _muni_row(
            SEOUL,
            "TC",
            C,
            hazard_intensity=53.8,
            hazard_intensity_unit="m/s",
            impact_function_id=9,
            calculation_status="NO_EXPOSURE_DATA",
            status_detail="no asset value supplied",
        ),
        _muni_row(
            SEOUL,
            "HW",
            K,
            hazard_intensity=37.6,
            hazard_intensity_unit="degC",
            calculation_status="HAZARD_ONLY",
        ),
        # Korea-local flood stays not implemented
        _muni_row(
            SEOUL,
            "RF",
            K,
            calculation_status="NOT_IMPLEMENTED",
            hazard_dataset=None,
            spatial_resolution=None,
            spatial_resolution_unit=None,
            spatial_unit_type=None,
        ),
    ]


def test_screening_rows_keep_hazard_and_null_money_never_zero() -> None:
    rows = _run_rows()
    seoul_rf = next(
        r for r in rows if r.facility_id == SEOUL and r.hazard_type == "RF" and r.model_id == C
    )
    assert seoul_rf.calculation_status == CalculationStatus.NO_EXPOSURE_DATA.value
    assert seoul_rf.hazard_intensity == 6.97 and seoul_rf.spatial_resolution == 150.0
    assert (
        seoul_rf.eal_usd is None
        and seoul_rf.potential_loss_usd is None
        and seoul_rf.risk_level is None
    )
    assert seoul_rf.asset_value_usd is None
    busan_rf = next(r for r in rows if r.facility_id == BUSAN and r.hazard_type == "RF")
    assert busan_rf.eal_usd == 0.0 and busan_rf.risk_level == "Low"  # a computed zero is priced


def test_municipality_summary_and_matrix_frames() -> None:
    rows = _run_rows()
    summary = rt.municipality_summary_frame(rows, dc.recommended_models())
    by = {line["Municipality ID"]: line for line in summary}
    assert by[GWANGJU]["Flood Risk"] == "High" and by[GWANGJU]["Flood EAL"] == pytest.approx(
        597_066.6
    )
    assert by[GWANGJU]["Financially Assessed"] == "Yes"
    assert by[GWANGJU]["Representative Point"] == "Boundary interior point"
    assert (
        by[GWANGJU]["Heatwave Status"] == "Hazard only"
        and by[GWANGJU]["Heatwave Tmax p95 (°C)"] == 37.8
    )
    assert by[BUSAN]["Flood EAL"] == 0 and by[BUSAN]["Flood Risk"] == "Low"
    assert by[SEOUL]["Flood Status"] == "No asset value" and by[SEOUL]["Flood EAL"] is None
    assert by[SEOUL]["Flood Risk"] is None and by[SEOUL]["Asset Value"] is None
    assert by[SEOUL]["Financially Assessed"] == "No — asset value not supplied"
    assert by[SEOUL]["Flood Hazard Intensity"] == "6.97 m"
    matrix = rt.municipality_risk_matrix_frame(rows, dc.recommended_models())
    bym = {line["Municipality ID"]: line for line in matrix}
    assert bym[GWANGJU] == {
        "Municipality ID": GWANGJU,
        "Municipality": "광주광역시",
        "Flood": "High",
        "Tropical Cyclone": "Low",
        "Heatwave": "Hazard only",
        "Calculated": "2 of 3 calculated",
    }
    assert (
        bym[SEOUL]["Flood"] == "No asset value" and bym[SEOUL]["Calculated"] == "0 of 3 calculated"
    )
    counts = rt.summary_counts(rows, dc.recommended_models())
    assert (
        counts["assets_analyzed"] == 3 and counts["high_risk"] == 1 and counts["hazard_only"] == 3
    )
    assert counts["not_available"] == 2  # Seoul's two screening rows — not zeros


def test_spatial_resolution_frame_lists_run_cells_then_extended_hazards() -> None:
    frame = rt.spatial_resolution_frame(_run_rows())
    labels = [(f["Hazard"], f["Model ID"]) for f in frame]
    assert ("Flood", C) in labels and ("Heatwave", K) in labels and ("Flood", K) in labels
    assert labels[-2:] == [("Drought", None), ("Sea-level rise", None)]
    flood = next(f for f in frame if f["Hazard"] == "Flood" and f["Model ID"] == C)
    assert flood["Spatial Resolution"] == 150.0 and flood["Unit"] == "arcsec"
    assert (
        flood["Approx. Korea Scale"] == "~4–5 km" and flood["Declared (registry)"] == "150 arcsec"
    )
    heat = next(f for f in frame if f["Hazard"] == "Heatwave")
    assert heat["Spatial Resolution"] == 0.05 and heat["Spatial Unit"] == "grid"
    local = next(f for f in frame if f["Hazard"] == "Flood" and f["Model ID"] == K)
    assert local["Spatial Resolution"] is None and local["Status"] == "NOT_IMPLEMENTED"
    assert all(f["Note"] == cov.RESOLUTION_NOTE for f in frame)


def test_hazard_results_columns_include_target_and_resolution() -> None:
    for col in (
        "assessment_target",
        "municipality_id",
        "municipality_name",
        "spatial_resolution",
        "spatial_resolution_unit",
        "spatial_unit_type",
    ):
        assert col in rt.HAZARD_RESULTS_COLUMNS
    hr = rt.hazard_results_frame(_run_rows())
    assert (
        hr[0]["assessment_target"] == "MUNICIPALITY" and hr[0]["municipality_name"] == "광주광역시"
    )


def _output(rows: list[ResultRow], target: str = "MUNICIPALITY") -> dict:  # type: ignore[type-arg]
    return {
        "status": "ok",
        "assessment_target": target,
        "municipality_dataset": mu.dataset_summary() if target == "MUNICIPALITY" else None,
        "climate_scenario": "rcp85",
        "target_year": 2040,
        "country": "KOR",
        "rows": [r.to_dict() for r in rows],
        "adapters": [
            {
                "hazard_type": "RF",
                "model_id": C,
                "spatial_resolution": {"value": 150.0, "unit": "arcsec", "unit_type": "grid cell"},
            },
        ],
        "readiness": {"summary": {}},
        "impact_function_fixed": {"RF": True},
        "detail": None,
    }


def test_municipality_workbook_has_the_required_sheets_and_blank_unpriced_cells() -> None:
    out = _output(_run_rows())
    frames = xl.frames_from_output(out)
    assert list(frames)[:5] == [
        "municipality_summary",
        "municipality_risk_matrix",
        "hazard_results",
        "spatial_resolution",
        "hazard_coverage",
    ]
    assert "global_vs_country" not in frames  # no Global run → no Global sheet
    buf = io.BytesIO()
    xl.write_workbook(frames, buf)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    titles = wb.sheetnames
    required = [xl.SHEET_TITLES[k] for k in xl.REQUIRED_MUNICIPALITY_SHEETS]
    assert [t for t in titles if t in required] == required
    ws = wb["Municipality Summary"]
    header = [c.value for c in ws[1]]
    seoul = next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == SEOUL)
    line = dict(zip(header, seoul, strict=True))
    assert line["Flood EAL"] is None and line["Flood Risk"] is None  # empty, not 0
    assert line["Flood Status"] == "No asset value"
    busan = dict(
        zip(
            header,
            next(r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == BUSAN),
            strict=True,
        )
    )
    assert busan["Flood EAL"] == 0  # a computed zero is written as 0
    cov_ws = wb["Hazard Coverage"]
    cov_rows = list(cov_ws.iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in cov_rows] == [
        "Flood",
        "Tropical Cyclone",
        "Heatwave",
        "Drought",
        "Sea-level rise",
    ]
    res_ws = wb["Spatial Resolution"]
    assert res_ws.max_row >= 1 + 3 + 2
    info = {r[0]: r[1] for r in wb["Run Info"].iter_rows(min_row=2, values_only=True)}
    assert info["assessment_target"] == "MUNICIPALITY" and info["n_municipalities"] == 3
    assert info["representative_point_warning"] == mu.REPRESENTATIVE_POINT_WARNING
    meth = {r[0]: r[1] for r in wb["Methodology"].iter_rows(min_row=2, values_only=True)}
    assert meth["Representative point"] == mu.REPRESENTATIVE_POINT_WARNING
    assert meth["Spatial resolution"] == cov.RESOLUTION_NOTE


def test_report_filename_by_target() -> None:
    out = _output(_run_rows())
    assert xl.report_filename(out, "2026-09-27T10:00:00") == (
        "Municipality_Physical_Risk_Report_3_Municipalities_20260927.xlsx"
    )
    fac = _output([], "FACILITY")
    assert (
        xl.report_filename(fac, "2026-09-27T10:00:00")
        == "Physical_Risk_Report_0_Assets_20260927.xlsx"
    )


def test_facility_workbook_is_unchanged_by_v02() -> None:
    row = ResultRow(
        facility_id="F1",
        facility_name="F1",
        hazard_type="RF",
        model_id=C,
        asset_value_usd=1e8,
        eal_usd=1000.0,
        potential_loss_usd=5e4,
        return_period_years=100.0,
        impact_function_id=22,
        calculation_status="FULL",
        requested_scenario="rcp85",
        served_scenario="rcp85",
    ).finalise()
    frames = xl.frames_from_output(_output([row], "FACILITY"))
    assert list(frames)[:3] == ["portfolio_summary", "asset_risk_matrix", "hazard_results"]
    assert "municipality_summary" not in frames and "spatial_resolution" not in frames
    assert frames["hazard_results"][0]["assessment_target"] == "FACILITY"


# --------------------------------------------------------------------------- #
# API                                                                           #
# --------------------------------------------------------------------------- #
def test_municipalities_endpoint_and_readiness_bundle(client) -> None:  # type: ignore[no-untyped-def]
    r = client.get("/api/libraries/municipalities")
    assert r.status_code == 200
    body = r.json()
    assert body["summary"]["total"] == len(mu.load_municipalities())
    assert {m["municipality_id"] for m in body["municipalities"]} >= {GWANGJU, SEOUL}
    assert client.get("/api/libraries/municipalities?level=METROPOLITAN").json()["municipalities"]
    assert client.get("/api/libraries/municipalities?level=PLANET").status_code == 400
    ready = client.get("/api/libraries/physical-risk-models").json()
    assert ready["display"]["coverage_table"][3]["Status"] == "NOT_READY"
    assert ready["display"]["municipality"]["warning"] == mu.REPRESENTATIVE_POINT_WARNING


def test_submit_validates_municipality_ids_and_values(client) -> None:  # type: ignore[no-untyped-def]
    session = client.post("/api/session").json()
    sid = session["id"]
    url = f"/api/session/{sid}/physical-risk-models"
    assert client.post(url, json={"assessment_target": "PLANET"}).status_code == 400
    assert client.post(url, json={"assessment_target": "MUNICIPALITY"}).status_code == 400
    r = client.post(
        url, json={"assessment_target": "MUNICIPALITY", "municipality_ids": ["SGIS-00000"]}
    )
    assert r.status_code == 400 and "unknown municipality" in r.json()["detail"]
    r = client.post(
        url,
        json={
            "assessment_target": "MUNICIPALITY",
            "municipality_ids": [GWANGJU],
            "asset_values": {SEOUL: 1.0},
        },
    )
    assert r.status_code == 400 and "not selected" in r.json()["detail"]
    r = client.post(
        url,
        json={
            "assessment_target": "MUNICIPALITY",
            "municipality_ids": [GWANGJU],
            "asset_values": {GWANGJU: 0},
        },
    )
    assert r.status_code == 400 and "positive" in r.json()["detail"]
