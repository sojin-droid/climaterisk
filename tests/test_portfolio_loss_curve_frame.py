"""The Portfolio Loss Curve sheet renders the worker's curves without computing anything."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from climaterisk.physical_risk.excel_export import frames_from_output, write_workbook
from climaterisk.physical_risk.metrics import ModelId
from climaterisk.physical_risk.results_table import export_frames, portfolio_loss_curve_frame

COUNTRY = ModelId.DATA_API_COUNTRY.value
GLOBAL = ModelId.GLOBAL_BASELINE.value


def _curve(level: str, name: str | None, model: str, **extra: object) -> dict[str, object]:
    losses = [
        {"return_period_years": 10.0, "loss": None, "status": "SHORTER_THAN_CURVE"},
        {"return_period_years": 25.0, "loss": 10.0, "status": "RESOLVED"},
        {"return_period_years": 50.0, "loss": 80.0, "status": "RESOLVED"},
        {
            "return_period_years": 100.0,
            "loss": 100.0,
            "status": "RESOLVED",
            "sum_of_standalone_losses": 180.0,
            "diversification_pct": 44.4,
        },
        {"return_period_years": 250.0, "loss": None, "status": "RETURN_PERIOD_NOT_RESOLVABLE"},
    ]
    return {
        "hazard_type": "RF",
        "model_id": model,
        "requested_scenario": "rcp45",
        "group_level": level,
        "group_name": name,
        "points_assessed": 3,
        "points_priced": 2,
        "priced_asset_value": 3000.0,
        "status": "OK",
        "eal": 2.3,
        "curve_return_period_min": 25.0,
        "curve_return_period_max": 100.0,
        "losses": losses,
        "asset_classes": {"commercial": 1, "industrial": 1},
        "impact_functions": [
            {
                "impact_function_id": 22,
                "impact_function_name": "Flood Asia JRC Commercial",
                "assets": 1,
            },
            {
                "impact_function_id": 23,
                "impact_function_name": "Flood Asia JRC Industrial",
                "assets": 1,
            },
        ],
        **extra,
    }


def test_frame_orders_portfolio_first_and_recommended_scope_first() -> None:
    curves = [
        _curve("PROVINCE", "부산광역시", COUNTRY),
        _curve("ASSET_CLASS", "residential", COUNTRY),
        _curve("PORTFOLIO", None, GLOBAL),
        _curve("ASSET_CLASS", "commercial", COUNTRY),
        _curve("PORTFOLIO", None, COUNTRY),
    ]
    lines = portfolio_loss_curve_frame(curves, preferred={"RF": COUNTRY})
    assert [(x["Level"], x["Group"], x["Recommended"]) for x in lines] == [
        ("Portfolio", "Portfolio", "Yes"),
        ("Portfolio", "Portfolio", "No"),
        ("Asset class", "Commercial", "Yes"),
        ("Asset class", "Residential", "Yes"),
        ("Province", "부산광역시", "Yes"),
    ]


def test_mixed_impact_functions_and_asset_classes_are_spelled_out() -> None:
    (line,) = portfolio_loss_curve_frame([_curve("PORTFOLIO", None, COUNTRY)])
    assert line["Asset Classes"] == "Commercial 1 · Industrial 1"
    assert line["Impact Functions"] == (
        "Flood Asia JRC Commercial (id 22) ×1 · Flood Asia JRC Industrial (id 23) ×1"
    )


def test_unresolvable_return_periods_are_blank_with_a_reason() -> None:
    (line,) = portfolio_loss_curve_frame([_curve("PORTFOLIO", None, COUNTRY)])
    assert line["10-yr Loss"] is None and line["250-yr Loss"] is None
    assert line["100-yr Loss"] == 100.0
    assert line["Sum of Standalone 100-yr Losses"] == 180.0
    assert line["Diversification Benefit % (100-yr)"] == 44.4
    assert "250-yr beyond" in line["Note"] and "10-yr shorter" in line["Note"]
    assert line["Curve Range (years)"] == "25 – 100"


def test_baseline_curves_follow_the_future_ones() -> None:
    lines = portfolio_loss_curve_frame(
        [_curve("PORTFOLIO", None, COUNTRY)],
        baseline_curves=[_curve("PORTFOLIO", None, COUNTRY, requested_scenario="historical")],
    )
    assert [(x["Run"], x["Scenario"]) for x in lines] == [
        ("Future", "rcp45"),
        ("Baseline", "historical"),
    ]


def test_sheet_only_when_curves_exist() -> None:
    assert "portfolio_loss_curve" not in export_frames([])
    frames = export_frames([], loss_curves=[_curve("PORTFOLIO", None, COUNTRY)])
    assert list(frames)[:2] == ["portfolio_totals", "portfolio_loss_curve"]


def test_workbook_writes_the_sheet_second_with_money_formats() -> None:
    output = {
        "assessment_target": "FACILITY",
        "rows": [],
        "portfolio_loss_curves": [_curve("PORTFOLIO", None, COUNTRY)],
    }
    buf = io.BytesIO()
    write_workbook(frames_from_output(output), buf)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    assert wb.sheetnames[:2] == ["Portfolio Totals", "Portfolio Loss Curve"]
    ws = wb["Portfolio Loss Curve"]
    header = [c.value for c in ws[1]]
    loss = ws.cell(row=2, column=header.index("100-yr Loss") + 1)
    assert loss.value == 100.0 and loss.number_format == '"$"#,##0'
    blank = ws.cell(row=2, column=header.index("250-yr Loss") + 1)
    assert blank.value is None  # never written as 0
