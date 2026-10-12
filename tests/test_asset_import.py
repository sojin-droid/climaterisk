"""Asset import templates — parsing, municipality joins, merge modes, routes.

CLIMADA-free. The rules under test are the ones that keep an import honest: a located
asset needs a value, a municipality never gets one invented, an ambiguous name is an
error, any bad row refuses the whole file with its line number, and a re-upload under
``merge`` replaces the same ids instead of duplicating them.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

from climaterisk.core.entities import Asset
from climaterisk.physical_risk import asset_import as ai
from climaterisk.physical_risk.municipalities import by_id

CHUNCHEON, SEOUL, GWANGJU = "SGIS-32010", "SGIS-11", "SGIS-24"


def _xlsx(header: list[str], rows: list[list[object]], sheet: str = "assets") -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(text: str) -> bytes:
    return text.encode("utf-8")


# --------------------------------------------------------------------------- #
# Templates                                                                    #
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("kind", ai.TEMPLATE_KINDS)
def test_every_template_downloads_and_round_trips_through_the_parser(kind: str) -> None:
    wb = load_workbook(io.BytesIO(ai.template_workbook(kind)))
    assert wb.sheetnames == ["README", "assets"]
    header = [c.value for c in wb["assets"][1]]
    assert header == list(ai.TEMPLATES[kind]["columns"])
    # the example rows are themselves a valid upload
    assets, report = ai.parse_assets(f"t.{kind}.xlsx", ai.template_workbook(kind), kind)
    assert len(assets) == len(ai.TEMPLATES[kind]["examples"]) == report.rows_read
    assert report.added == [a.id for a in assets]


def test_unknown_template_kind_is_refused() -> None:
    with pytest.raises(KeyError):
        ai.template_workbook("nuts2")
    with pytest.raises(ai.AssetImportError, match="unknown template kind"):
        ai.parse_assets("x.csv", _csv("facility_id,lat,lon,value\nA,1,2,3\n"), "nuts2")


# --------------------------------------------------------------------------- #
# latlon                                                                       #
# --------------------------------------------------------------------------- #
def test_latlon_csv_with_aliases_and_a_sector_or_property_type() -> None:
    text = (
        "id,name,lat,lon,value,currency,type,people\n"
        "A,Alpha,37.5,127.0,1000000,KRW,steel,40\n"
        "B,,35.1,129.0,2e6,,office,\n"
    )
    assets, report = ai.parse_assets("portfolio.csv", _csv(text), "latlon")
    a, b = assets
    assert a.id == "A" and a.name == "Alpha" and a.sector.value == "steel" and a.headcount == 40
    assert a.currency == "KRW" and a.value == 1_000_000.0 and "property_type" not in a.properties
    assert b.name == "B" and b.sector.value == "real_estate"  # default sector, label carried
    assert b.properties["property_type"] == "office" and b.currency == "USD" and b.headcount is None
    assert report.rows_read == 2 and report.added == ["A", "B"]


def test_latlon_refuses_the_whole_file_with_line_numbers() -> None:
    text = (
        "facility_id,latitude,longitude,asset_value_usd\n"
        "A,37.5,127.0,100\n"
        "A,137.5,127.0,100\n"
        ",37.5,127.0,100\n"
        "C,37.5,127.0,\n"
        "D,37.5,127.0,-5\n"
        "E,north,127.0,1\n"
    )
    with pytest.raises(ai.AssetImportError) as exc:
        ai.parse_assets("p.csv", _csv(text), "latlon")
    msgs = exc.value.errors
    assert any(m.startswith("line 3:") and "out of range" in m for m in msgs)
    assert any(m.startswith("line 4:") and "empty facility_id" in m for m in msgs)
    assert any(m.startswith("line 5:") and "required for a located asset" in m for m in msgs)
    assert any(m.startswith("line 6:") and "negative" in m for m in msgs)
    assert any(m.startswith("line 7:") and "not a number" in m for m in msgs)


def test_latlon_needs_its_required_columns() -> None:
    with pytest.raises(ai.AssetImportError, match="missing required column"):
        ai.parse_assets("p.csv", _csv("facility_id,lat\nA,1\n"), "latlon")
    with pytest.raises(ai.AssetImportError, match="no header"):
        ai.parse_assets("p.csv", b"", "latlon")
    with pytest.raises(ai.AssetImportError, match="no data rows"):
        ai.parse_assets("p.csv", _csv("facility_id,lat,lon,value\n"), "latlon")


def test_xlsx_is_read_from_the_assets_sheet_or_the_first_sheet() -> None:
    header = ["facility_id", "latitude", "longitude", "asset_value_usd"]
    rows: list[list[object]] = [["A", 37.5, 127.0, 100.0]]
    assets, _ = ai.parse_assets("p.xlsx", _xlsx(header, rows), "latlon")
    assert assets[0].lat == 37.5
    assets, _ = ai.parse_assets("p.xlsx", _xlsx(header, rows, sheet="Sheet1"), "latlon")
    assert assets[0].value == 100.0
    assets, _ = ai.parse_assets("", _xlsx(header, rows), "latlon")  # sniffed from the bytes
    assert assets[0].id == "A"


# --------------------------------------------------------------------------- #
# sigungu / sido                                                               #
# --------------------------------------------------------------------------- #
def test_sigungu_rows_land_on_the_official_representative_point() -> None:
    text = (
        "municipality_id,municipality_name,province_name,asset_value_usd,property_type,facility_id\n"
        f"{CHUNCHEON},,,50000000,office,\n"
        ",중구,서울특별시,,,\n"
    )
    assets, report = ai.parse_assets("m.csv", _csv(text), "sigungu")
    chun, jung = assets
    m = by_id([CHUNCHEON])[0]
    assert chun.id == CHUNCHEON and chun.name == m.municipality_name
    assert (chun.lat, chun.lon) == (m.latitude, m.longitude)
    assert chun.value == 50_000_000.0 and chun.properties["property_type"] == "office"
    assert (
        chun.properties["point_type"] == m.point_type
        and chun.properties["province_name"] == m.province_name
    )
    assert jung.name == "중구" and jung.properties["province_name"] == "서울특별시"
    assert jung.value == 0.0  # no value supplied -> screening only, nothing invented
    assert report.added == [CHUNCHEON, jung.id] and jung.id.startswith("SGIS-")


def test_sigungu_ambiguous_or_unknown_names_are_errors_not_guesses() -> None:
    text = "municipality_name\n중구\n없는시\n"
    with pytest.raises(ai.AssetImportError) as exc:
        ai.parse_assets("m.csv", _csv(text), "sigungu")
    msgs = exc.value.errors
    assert any("line 2:" in m and "ambiguous" in m and "province_name" in m for m in msgs)
    assert any("line 3:" in m and "no sigungu named" in m for m in msgs)
    with pytest.raises(ai.AssetImportError, match="unknown sigungu municipality_id"):
        ai.parse_assets("m.csv", _csv("municipality_id\nSGIS-00000\n"), "sigungu")
    with pytest.raises(ai.AssetImportError, match="needs municipality_id or municipality_name"):
        ai.parse_assets("m.csv", _csv("asset_value_usd\n1\n"), "sigungu")


def test_sido_rows_use_the_province_unit_and_refuse_a_sigungu_id() -> None:
    assets, _ = ai.parse_assets("s.csv", _csv(f"municipality_id\n{SEOUL}\n"), "sido")
    assert assets[0].properties["municipality_level"] == "METROPOLITAN"
    assets, _ = ai.parse_assets("s.csv", _csv("municipality_name,value\n부산광역시,1\n"), "sido")
    assert assets[0].value == 1.0 and assets[0].name == "부산광역시"
    with pytest.raises(ai.AssetImportError, match="unknown sido municipality_id"):
        ai.parse_assets("s.csv", _csv(f"municipality_id\n{CHUNCHEON}\n"), "sido")


# --------------------------------------------------------------------------- #
# Merge modes                                                                  #
# --------------------------------------------------------------------------- #
def _asset(aid: str, value: float) -> Asset:
    return Asset(id=aid, name=aid, lat=37.0, lon=127.0, value=value)


def test_merge_replaces_same_ids_in_place_and_keeps_the_rest() -> None:
    existing = [_asset("A", 1.0), _asset("B", 2.0)]
    merged, removed = ai.merge_assets(existing, [_asset("B", 20.0), _asset("C", 3.0)], "merge")
    assert [a.id for a in merged] == ["A", "B", "C"] and merged[1].value == 20.0 and removed == []
    replaced, removed = ai.merge_assets(existing, [_asset("B", 20.0)], "replace")
    assert [a.id for a in replaced] == ["B"] and removed == ["A"]
    with pytest.raises(ValueError, match="mode"):
        ai.merge_assets(existing, [], "upsert")


# --------------------------------------------------------------------------- #
# Routes                                                                       #
# --------------------------------------------------------------------------- #
def test_template_routes_list_and_download(client) -> None:  # type: ignore[no-untyped-def]
    listing = client.get("/api/libraries/asset-templates").json()["templates"]
    assert [t["kind"] for t in listing] == list(ai.TEMPLATE_KINDS)
    r = client.get("/api/libraries/asset-templates/sigungu.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert "climaterisk_assets_sigungu.xlsx" in r.headers["content-disposition"]
    assert client.get("/api/libraries/asset-templates/nuts2.xlsx").status_code == 404


def test_import_route_merges_reuploads_and_refuses_bad_files_whole(client) -> None:  # type: ignore[no-untyped-def]
    sid = client.post("/api/session").json()["id"]
    url = f"/api/session/{sid}/assets/import"
    good = _csv("facility_id,lat,lon,value\nA,37.5,127.0,100\nB,35.1,129.0,200\n")
    r = client.post(url, files={"file": ("p.csv", good, "text/csv")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert [a["id"] for a in body["portfolio"]["assets"]] == ["A", "B"]
    assert body["report"]["added"] == ["A", "B"] and body["report"]["mode"] == "merge"
    # the session really holds them
    assert len(client.get(f"/api/session/{sid}").json()["assets"]) == 2
    # re-upload with one changed value and one new row: same ids replaced, not duplicated
    again = _csv("facility_id,lat,lon,value\nB,35.1,129.0,250\nC,36.0,128.0,300\n")
    body = client.post(url, files={"file": ("p.csv", again, "text/csv")}).json()
    assets = {a["id"]: a for a in body["portfolio"]["assets"]}
    assert list(assets) == ["A", "B", "C"] and assets["B"]["value"] == 250.0
    assert body["report"]["replaced"] == ["B"] and body["report"]["added"] == ["C"]
    # replace mode drops what the file does not carry, and says so
    body = client.post(f"{url}?mode=replace", files={"file": ("p.csv", again, "text/csv")}).json()
    assert [a["id"] for a in body["portfolio"]["assets"]] == ["B", "C"]
    assert body["report"]["removed"] == ["A"]
    # a bad row refuses the whole file: nothing changes
    bad = _csv("facility_id,lat,lon,value\nD,37.5,127.0,1\nE,999,127.0,1\n")
    r = client.post(url, files={"file": ("p.csv", bad, "text/csv")})
    assert r.status_code == 422 and any("line 3" in e for e in r.json()["detail"]["errors"])
    assert [a["id"] for a in client.get(f"/api/session/{sid}").json()["assets"]] == ["B", "C"]
    # parameter and session validation
    assert client.post(f"{url}?kind=nuts2", files={"file": ("p.csv", good)}).status_code == 400
    assert client.post(f"{url}?mode=upsert", files={"file": ("p.csv", good)}).status_code == 400
    assert client.post(url, files={"file": ("p.csv", b"")}).status_code == 400
    assert (
        client.post("/api/session/nope/assets/import", files={"file": ("p.csv", good)}).status_code
        == 404
    )


def test_import_route_places_municipalities_for_the_models_tab(client) -> None:  # type: ignore[no-untyped-def]
    sid = client.post("/api/session").json()["id"]
    text = f"municipality_id,asset_value_usd\n{GWANGJU},80000000\n"
    r = client.post(
        f"/api/session/{sid}/assets/import?kind=sido",
        files={"file": ("sido.csv", _csv(text), "text/csv")},
    )
    assert r.status_code == 200, r.text
    (asset,) = r.json()["portfolio"]["assets"]
    m = by_id([GWANGJU])[0]
    assert asset["id"] == GWANGJU and asset["lat"] == m.latitude and asset["value"] == 80_000_000.0
    assert asset["properties"]["municipality_level"] == "METROPOLITAN"
