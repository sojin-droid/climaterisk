"""Phase 5 — batch table to Excel, same-model climate multipliers, CSV facilities.

CLIMADA-free. The workbook is read back with openpyxl and checked for the rules that make
it trustworthy: the required sheets exist in order, every result row carries ``model_id``
and ``calculation_status``, an unpriced field is an **empty cell** (never ``0``), heat rows
have no financial numbers, and the comparison columns are the ones the brief lists.
"""

from __future__ import annotations

import importlib.util
import io
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
if str(REPO / "src") not in sys.path:
    sys.path.insert(0, str(REPO / "src"))

from climaterisk.physical_risk import excel_export as xl  # noqa: E402
from climaterisk.physical_risk.metrics import (  # noqa: E402
    CalculationStatus,
    ModelId,
    ResultRow,
    attach_climate_multipliers,
)
from climaterisk.physical_risk.results_table import HAZARD_RESULTS_COLUMNS  # noqa: E402

G, C, K = ModelId.GLOBAL_BASELINE.value, ModelId.DATA_API_COUNTRY.value, ModelId.KOREA_LOCAL.value

SPEC_ROW_FIELDS = (
    "risk_level",
    "risk_level_criteria",
    "potential_loss_usd",
    "potential_loss_return_period_years",
    "eal_usd",
    "eal_as_pct_of_assets",
    "probability",
    "return_period_years",
    "climate_change_multiplier",
    "calculation_status",
    "hazard_source",
    "impact_function_id",
)
SPEC_COMPARISON_FIELDS = (
    "eal_change_usd",
    "eal_change_pct",
    "eal_as_pct_assets_change_pp",
    "potential_loss_change_usd",
    "potential_loss_change_pct",
    "hazard_intensity_change_pct",
)


def _row(fid: str, haz: str, model: str, **kw) -> ResultRow:  # type: ignore[no-untyped-def]
    base = {
        "facility_id": fid,
        "facility_name": f"Facility {fid}",
        "hazard_type": haz,
        "model_id": model,
        "asset_value_usd": 100_000_000.0,
        "asset_value_currency": "USD",
        "scenario": "rcp60",
        "requested_scenario": "rcp60",
        "served_scenario": "rcp60",
        "time_horizon": "2030–2050",
        "hazard_source": f"synthetic {model}",
    }
    base.update(kw)
    return ResultRow(**base).finalise()  # type: ignore[arg-type]


def _priced(fid: str, haz: str, model: str, eal: float, pml: float, **kw) -> ResultRow:  # type: ignore[no-untyped-def]
    return _row(
        fid,
        haz,
        model,
        eal_usd=eal,
        potential_loss_usd=pml,
        potential_loss_return_period_years=100.0,
        return_period_years=100.0,
        impact_function_id=22 if haz == "RF" else 9,
        hazard_intensity=1.2 if haz == "RF" else 50.0,
        hazard_intensity_unit="m" if haz == "RF" else "m/s",
        calculation_status=CalculationStatus.FULL.value,
        **kw,
    )


def _output(n_facilities: int = 2) -> dict:  # type: ignore[type-arg]
    rows: list[ResultRow] = []
    for i in range(n_facilities):
        fid = f"F{i + 1}"
        rows += [
            _priced(fid, "RF", G, 75_000.0 * (i + 1), 7_200_000.0),
            _priced(fid, "RF", C, 90_000.0 * (i + 1), 8_600_000.0),
            _row(fid, "RF", K, calculation_status="NOT_IMPLEMENTED", status_detail="no adapter"),
            _priced(fid, "TC", G, 3_576.18, 121_636.62),
            _priced(fid, "TC", C, 3_576.18, 121_636.62),
            _row(fid, "TC", K, calculation_status="NOT_IMPLEMENTED"),
            _row(fid, "HW", G, calculation_status="NO_HAZARD_DATA"),
            _row(fid, "HW", C, calculation_status="NO_HAZARD_DATA"),
            _row(
                fid,
                "HW",
                K,
                hazard_intensity=37.1,
                hazard_intensity_unit="degC",
                hazard_source="KMA 남한상세 TAMAX via local catalog",
                calculation_status=CalculationStatus.HAZARD_ONLY.value,
            ),
        ]
    from climaterisk.physical_risk.models import readiness

    return {
        "status": "ok",
        "climate_scenario": "rcp60",
        "target_year": 2050,
        "country": "KOR",
        "rows": [r.to_dict() for r in rows],
        "readiness": readiness(),
        "impact_function_fixed": {"RF": True, "TC": True, "HW": True},
        "detail": None,
    }


def _book(output: dict):  # type: ignore[type-arg,no-untyped-def]
    return load_workbook(io.BytesIO(xl.workbook_bytes(output)))


def _sheet_rows(ws) -> list[dict]:  # type: ignore[no-untyped-def,type-arg]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    return [dict(zip(header, r, strict=True)) for r in it]


# --------------------------------------------------------------------------- #
# Sheets and columns                                                           #
# --------------------------------------------------------------------------- #
def test_workbook_has_the_required_sheets_in_order_and_the_full_result_columns() -> None:
    wb = _book(_output())
    names = wb.sheetnames
    titles = [xl.SHEET_TITLES[k] for k in xl.SHEET_ORDER]
    assert [n for n in names if n in titles] == names  # only known sheets
    assert names == [t for t in titles if t in names]  # in the canonical order
    assert {xl.SHEET_TITLES[k] for k in xl.REQUIRED_SHEETS} <= set(names)
    assert "Climate Change" not in names  # no baseline scenario was run
    header = [c.value for c in wb["Hazard Results"][1]]
    assert header == list(HAZARD_RESULTS_COLUMNS)
    assert set(SPEC_ROW_FIELDS) <= set(header)
    assert {"model_id", "calculation_status"} <= set(header)


def test_every_result_row_keeps_model_id_and_status_and_unpriced_cells_are_empty() -> None:
    out = _output()
    rows = _sheet_rows(_book(out)["Hazard Results"])
    assert len(rows) == len(out["rows"]) == 18
    for r in rows:
        assert r["model_id"] in {G, C, K} and r["calculation_status"]
    heat = [r for r in rows if r["hazard_type"] == "HW"]
    assert len(heat) == 6
    for r in heat:
        for f in ("eal_usd", "potential_loss_usd", "eal_as_pct_of_assets", "risk_level"):
            assert r[f] is None, f"{f} must be an empty cell, got {r[f]!r}"
    kma = [r for r in heat if r["model_id"] == K]
    assert all(
        r["calculation_status"] == "HAZARD_ONLY" and r["hazard_intensity"] == 37.1 for r in kma
    )
    korea_local = [r for r in rows if r["model_id"] == K and r["hazard_type"] != "HW"]
    assert all(
        r["calculation_status"] == "NOT_IMPLEMENTED" and r["eal_usd"] is None for r in korea_local
    )
    priced = [r for r in rows if r["calculation_status"] == "FULL"]
    assert all(
        r["eal_usd"] is not None and r["risk_level"] in {"Low", "Medium", "High"} for r in priced
    )
    # a real zero would be written as 0 — none of the unpriced rows became one
    assert not any(r["eal_usd"] == 0 for r in rows if r["calculation_status"] != "FULL")


def test_global_vs_country_sheet_carries_every_change_form_the_brief_lists() -> None:
    wb = _book(_output())
    cmp = _sheet_rows(wb["Global vs Country"])
    header = set(cmp[0])
    assert set(SPEC_COMPARISON_FIELDS) <= header
    assert {"baseline_model_id", "country_model_id", "same_impact_function", "available"} <= header
    by = {(r["facility_id"], r["hazard_type"]): r for r in cmp}
    rf = by[("F1", "RF")]
    assert rf["available"] is True and rf["same_impact_function"] is True
    assert rf["eal_change_usd"] == pytest.approx(15_000.0)
    assert rf["eal_change_pct"] == pytest.approx(20.0)
    assert rf["eal_as_pct_assets_change_pp"] == pytest.approx(0.015)
    assert rf["potential_loss_change_pct"] == pytest.approx(19.444444, rel=1e-6)
    hw = by[("F1", "HW")]
    assert hw["available"] is False and hw["eal_change_pct"] is None
    tc = by[("F1", "TC")]
    assert tc["eal_change_pct"] == pytest.approx(0.0) and tc[
        "hazard_intensity_change_pct"
    ] == pytest.approx(0.0)
    gvk = _sheet_rows(wb["Global vs Korea Local"])
    assert all(r["available"] is False for r in gvk)


def test_asset_summary_and_methodology_and_run_info_sheets() -> None:
    wb = _book(_output())
    from climaterisk.physical_risk.results_table import asset_summary_frame

    summary = {(r["facility_id"], r["model_id"]): r for r in asset_summary_frame(_output()["rows"])}
    assert summary[("F1", G)]["eal_usd"] == pytest.approx(75_000.0 + 3_576.18)
    assert summary[("F1", G)]["priced_hazards"] == "RF,TC"
    assert "HW:NO_HAZARD_DATA" in summary[("F1", G)]["unpriced_hazards"]
    assert summary[("F2", K)]["eal_usd"] is None
    meth = {r["key"]: r["value"] for r in _sheet_rows(wb["Methodology"])}
    assert "never a climate-change multiplier" in meth["Climate Change Multiplier"]
    assert "NOT an official GRESB threshold" in meth["Risk Level — attribution"]
    assert "Technical definitions — model" in meth  # the technical lines follow the plain ones
    assert next(iter(meth)) == "What was assessed"
    info = {r["key"]: r["value"] for r in _sheet_rows(wb["Run Info"])}
    assert info["requested_scenario"] == "rcp60" and info["n_facilities"] == 2
    assert info[f"readiness.RF.{K}"] == "NOT_IMPLEMENTED"
    assert "not treated as Korea-local" in info[f"definition.{C}"]


def test_writer_refuses_frames_without_the_identity_columns() -> None:
    with pytest.raises(ValueError, match="hazard_results"):
        xl.write_workbook({"methodology": []}, io.BytesIO())
    with pytest.raises(ValueError, match="model_id"):
        xl.write_workbook({"hazard_results": [{"facility_id": "x"}]}, io.BytesIO())


# --------------------------------------------------------------------------- #
# Same-model climate multipliers                                               #
# --------------------------------------------------------------------------- #
def test_attach_climate_multipliers_pairs_within_a_model_only() -> None:
    fut = {
        "scenario": "rcp45",
        "requested_scenario": "rcp45",
        "served_scenario": "rcp45",
        "time_horizon": "2040",
    }
    hist = {
        "scenario": "historical",
        "requested_scenario": "historical",
        "served_scenario": "historical",
        "time_horizon": "1980–2020",
    }
    fut_g = _priced("F1", "TC", G, 150.0, 1000.0, **fut)
    fut_c = _priced("F1", "TC", C, 300.0, 1000.0, **fut)
    heat = _row("F1", "HW", K, calculation_status="HAZARD_ONLY", **fut)
    base_g = _priced("F1", "TC", G, 100.0, 900.0, **hist)
    base_heat = _row("F1", "HW", K, calculation_status="HAZARD_ONLY", **hist)
    recs = attach_climate_multipliers([fut_g, fut_c, heat], [base_g, base_heat])
    by = {(r["hazard_type"], r["model_id"]): r for r in recs}
    assert fut_g.climate_change_multiplier == pytest.approx(1.5)
    assert by[("TC", G)]["climate_change_multiplier"] == pytest.approx(1.5)
    assert (
        fut_c.climate_change_multiplier is None
    )  # no DATA_API_COUNTRY baseline — never borrowed from G
    assert "no baseline row" in by[("TC", C)]["detail"]
    assert heat.climate_change_multiplier is None and "no EAL" in by[("HW", K)]["detail"]
    # and the workbook gets the climate_change sheet when multipliers exist
    out = _output(1)
    out["baseline_scenario"] = "historical"
    out["climate_change_multipliers"] = recs
    wb = _book(out)
    assert "Climate Change" in wb.sheetnames
    cc = _sheet_rows(wb["Climate Change"])
    assert {"model_id", "climate_change_multiplier", "detail"} <= set(cc[0])


# --------------------------------------------------------------------------- #
# Batch CLI — CSV parsing is pure                                              #
# --------------------------------------------------------------------------- #
def _cli():  # type: ignore[no-untyped-def]
    spec = importlib.util.spec_from_file_location(
        "physical_risk_batch", REPO / "scripts" / "physical_risk_batch.py"
    )
    assert spec is not None and spec.loader is not None
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_cli_reads_facilities_with_aliases_and_refuses_missing_columns(tmp_path: Path) -> None:
    cli = _cli()
    csv_path = tmp_path / "f.csv"
    csv_path.write_text(
        "facility_id,name,lat,lon,asset_value_usd,property_type\n"
        "KR-1,Concordian,37.5,127.0,100000000,office\n"
        "KR-2,Busan DC,35.1,129.0,,data_center\n",
        encoding="utf-8",
    )
    facs = cli.read_facilities(csv_path)
    assert [f["facility_id"] for f in facs] == ["KR-1", "KR-2"]
    assert facs[0]["facility_name"] == "Concordian" and facs[0]["latitude"] == 37.5
    assert facs[0]["asset_value_usd"] == 1e8 and facs[0]["asset_value_currency"] == "USD"
    assert (
        facs[1]["asset_value_usd"] is None
    )  # blank value -> None (NO_EXPOSURE_DATA downstream), not 0
    bad = tmp_path / "bad.csv"
    bad.write_text("facility_id,lat\nX,1\n", encoding="utf-8")
    with pytest.raises(ValueError, match="required column"):
        cli.read_facilities(bad)
    worse = tmp_path / "worse.csv"
    worse.write_text("facility_id,lat,lon,asset_value_usd\nX,north,127,1\n", encoding="utf-8")
    with pytest.raises(ValueError, match="line 2"):
        cli.read_facilities(worse)


def test_export_route_404s_for_an_unknown_run(client) -> None:  # type: ignore[no-untyped-def]
    r = client.get("/api/session/nope/run/nope/physical-risk-export.xlsx")
    assert r.status_code == 404
