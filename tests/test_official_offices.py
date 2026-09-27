"""V0.3 POC — official office points for the 7 metropolitan city halls.

CLIMADA-free. The rules under test:

* only government-published coordinates: an AVAILABLE row is ``GOVERNMENT_PUBLISHED`` with a
  named source, URL, date and polygon check; every other target carries **no** coordinates
  (never geocoded, never the representative point);
* the Busan row is exactly the published source row (text-identical coordinates);
* the V0.2 representative points are untouched and stay the default anchor; the office
  point is a separate facility (own id and ``point_type``) with the same user-supplied value
  or none — never an estimated one;
* Official Office is refused for an unsupported municipality (no silent fallback);
* the comparison frame shows both anchors, a descriptive distance and "Not available in
  V0.3 POC" where no official coordinate exists; a representative-only run's frames are
  unchanged (no "Coordinate Type", no comparison sheet).
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
if str(REPO / "src") not in sys.path:
    sys.path.insert(0, str(REPO / "src"))

from climaterisk.engines.base import PhysicalRiskModelsRequest  # noqa: E402
from climaterisk.physical_risk import excel_export as xl  # noqa: E402
from climaterisk.physical_risk import municipalities as mu  # noqa: E402
from climaterisk.physical_risk import official_offices as oo  # noqa: E402
from climaterisk.physical_risk import results_table as rt  # noqa: E402
from climaterisk.physical_risk.metrics import ModelId, ResultRow  # noqa: E402

BUSAN, SEOUL, GWANGJU = "SGIS-21", "SGIS-11", "SGIS-24"
C, K = ModelId.DATA_API_COUNTRY.value, ModelId.KOREA_LOCAL.value


# --------------------------------------------------------------------------- #
# Distance (descriptive) — analytical values                                    #
# --------------------------------------------------------------------------- #
def test_haversine_matches_analytical_distances() -> None:
    R = oo.EARTH_RADIUS_KM
    assert oo.haversine_km(10.0, 20.0, 10.0, 20.0) == 0.0
    # one degree of latitude along a meridian = R * pi / 180
    np_ = pytest.approx(R * math.pi / 180, rel=1e-12)
    assert oo.haversine_km(0.0, 30.0, 1.0, 30.0) == np_
    # one degree of longitude on the equator is the same arc
    assert oo.haversine_km(0.0, 0.0, 0.0, 1.0) == np_
    # antipodes = half the circumference
    assert oo.haversine_km(0.0, 0.0, 0.0, 180.0) == pytest.approx(math.pi * R, rel=1e-12)
    # symmetric
    a = oo.haversine_km(35.2, 129.05, 35.18, 129.08)
    assert a == pytest.approx(oo.haversine_km(35.18, 129.08, 35.2, 129.05), rel=1e-15)


# --------------------------------------------------------------------------- #
# Dataset                                                                       #
# --------------------------------------------------------------------------- #
def test_dataset_has_exactly_the_seven_targets_and_only_published_coordinates() -> None:
    offices = oo.load_offices()
    assert [o.municipality_id for o in offices] == list(oo.TARGETS)
    assert [o.office_name for o in offices] == list(oo.TARGETS.values())
    available = [o for o in offices if o.available]
    assert [o.municipality_id for o in available] == [BUSAN]
    for o in offices:
        if o.available:
            assert o.coordinate_method == oo.COORDINATE_METHOD == "GOVERNMENT_PUBLISHED"
            assert o.source_url and o.source_dataset and o.source_last_modified
            assert o.polygon_check in oo.POLYGON_CHECKS
        else:
            assert o.office_status == oo.STATUS_NOT_AVAILABLE
            assert o.latitude is None and o.longitude is None and o.coordinate_method is None
            assert o.note  # the reason is stated


def test_busan_row_is_the_published_source_row() -> None:
    b = oo.offices_by_id()[BUSAN]
    with oo.dataset_path().open(encoding="utf-8") as fh:
        raw = {r["municipality_id"]: r for r in csv.DictReader(fh)}[BUSAN]
    # the published text, unrounded
    assert (raw["latitude"], raw["longitude"]) == ("35.1799490", "129.0751049")
    assert b.source_dataset_id == "15025212"
    assert b.source_url == "https://www.data.go.kr/data/15025212/fileData.do"
    assert b.source_last_modified == "2026-05-11"
    assert b.source_row_name == "부산광역시청, 부산광역시의회"
    assert b.address == "부산광역시 연제구 중앙대로 1001"
    assert b.polygon_check == "INSIDE"
    rep = mu.by_id([BUSAN])[0]
    assert b.distance_to_representative_km == pytest.approx(
        oo.haversine_km(rep.latitude, rep.longitude, b.latitude, b.longitude), abs=5e-4
    )
    registry = json.loads(oo.registry_path().read_text(encoding="utf-8"))
    src = next(t for t in registry["targets"] if t["municipality_id"] == BUSAN)["source"]
    assert src["row_name"] == b.source_row_name and len(src["sha256"]) == 64


def _rows() -> list[dict[str, str]]:
    with oo.dataset_path().open(encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def test_validate_rejects_geocoded_substituted_or_incomplete_rows() -> None:
    assert oo.validate(_rows()) == []
    geocoded = _rows()
    geocoded[1]["coordinate_method"] = "ADDRESS_GEOCODING"
    assert any("GOVERNMENT_PUBLISHED" in p for p in oo.validate(geocoded))
    substituted = _rows()  # a not-available city given the representative point
    seoul = mu.by_id([SEOUL])[0]
    substituted[0]["latitude"], substituted[0]["longitude"] = (
        str(seoul.latitude),
        str(seoul.longitude),
    )
    assert any("must not carry coordinates" in p for p in oo.validate(substituted))
    assert any("targets without a row" in p for p in oo.validate(_rows()[:6]))
    renamed = _rows()
    renamed[1]["office_name"] = "부산시청"
    assert any("office_name must be" in p for p in oo.validate(renamed))
    no_source = _rows()
    no_source[1]["source_url"] = ""
    assert any("source_url must be stated" in p for p in oo.validate(no_source))
    extra = [*_rows(), {**_rows()[1], "municipality_id": "SGIS-29"}]
    assert any("not a V0.3 POC target" in p for p in oo.validate(extra))


def test_representative_points_are_untouched() -> None:
    ms = mu.load_municipalities()
    assert len(ms) == 269
    assert {m.point_type for m in ms} == {mu.POINT_TYPE_BOUNDARY_INTERIOR}
    b = mu.by_id([BUSAN])[0]
    assert (b.latitude, b.longitude) == (35.211169, 129.053955)
    assert b.office_name is None and b.address is None


# --------------------------------------------------------------------------- #
# Facilities and request                                                       #
# --------------------------------------------------------------------------- #
def test_office_facility_is_a_separate_point_with_no_invented_value() -> None:
    b, office = mu.by_id([BUSAN])[0], oo.offices_by_id()[BUSAN]
    f = oo.office_facility(b, office)
    assert f["facility_id"] == "SGIS-21#OFFICIAL_OFFICE" != b.municipality_id
    assert f["municipality_id"] == BUSAN and f["facility_name"] == "부산광역시청"
    assert (f["latitude"], f["longitude"]) == (office.latitude, office.longitude)
    assert f["point_type"] == oo.POINT_TYPE_OFFICIAL_OFFICE
    assert f["asset_value_usd"] is None
    assert oo.office_facility(b, office, 5e7)["asset_value_usd"] == 5e7  # user-supplied only
    with pytest.raises(ValueError, match="not available"):
        oo.office_facility(mu.by_id([SEOUL])[0], oo.offices_by_id()[SEOUL])


def test_request_anchors_default_to_the_v02_representative_point() -> None:
    ids = [BUSAN, SEOUL]
    default = PhysicalRiskModelsRequest.from_municipalities("s", "rcp85", ids)
    assert [f.facility_id for f in default.facilities] == ids
    assert {f.point_type for f in default.facilities} == {mu.POINT_TYPE_BOUNDARY_INTERIOR}
    both = PhysicalRiskModelsRequest.from_municipalities(
        "s", "rcp85", ids, asset_values={BUSAN: 1e6}, anchors=list(oo.ANCHORS)
    )
    assert [f.facility_id for f in both.facilities] == [BUSAN, "SGIS-21#OFFICIAL_OFFICE", SEOUL]
    assert [f.asset_value_usd for f in both.facilities] == [1e6, 1e6, None]
    office_only = PhysicalRiskModelsRequest.from_municipalities(
        "s", "rcp85", [BUSAN], anchors=[oo.ANCHOR_OFFICIAL_OFFICE]
    )
    assert [f.point_type for f in office_only.facilities] == [oo.POINT_TYPE_OFFICIAL_OFFICE]


# --------------------------------------------------------------------------- #
# Frames and workbook                                                          #
# --------------------------------------------------------------------------- #
def _row(fid: str, mid: str, haz: str, model: str, point_type: str, lat: float, lon: float, **kw):  # type: ignore[no-untyped-def]
    m = mu.by_id([mid])[0]
    base = {
        "facility_id": fid,
        "facility_name": m.municipality_name,
        "hazard_type": haz,
        "model_id": model,
        "assessment_target": "MUNICIPALITY",
        "municipality_id": mid,
        "municipality_name": m.municipality_name,
        "municipality_level": m.municipality_level,
        "province_name": m.province_name,
        "point_type": point_type,
        "latitude": lat,
        "longitude": lon,
        "requested_scenario": "rcp85",
        "served_scenario": "rcp85",
        "hazard_dataset": f"synthetic_{haz}",
    }
    base.update(kw)
    return ResultRow(**base).finalise()  # type: ignore[arg-type]


def _two_anchor_rows() -> list[ResultRow]:
    b, s = mu.by_id([BUSAN])[0], mu.by_id([SEOUL])[0]
    o = oo.offices_by_id()[BUSAN]
    assert o.latitude is not None and o.longitude is not None
    rp, op = mu.POINT_TYPE_BOUNDARY_INTERIOR, oo.POINT_TYPE_OFFICIAL_OFFICE
    ofid = oo.office_facility_id(BUSAN)
    rows = []
    for fid, pt, lat, lon, depth in (
        (BUSAN, rp, b.latitude, b.longitude, 0.0),
        (ofid, op, o.latitude, o.longitude, 1.5),
    ):
        rows.append(
            _row(
                fid,
                BUSAN,
                "RF",
                C,
                pt,
                lat,
                lon,
                calculation_status="NO_EXPOSURE_DATA",
                hazard_intensity=depth,
                hazard_intensity_unit="m",
            )
        )
        rows.append(
            _row(
                fid,
                BUSAN,
                "TC",
                C,
                pt,
                lat,
                lon,
                calculation_status="NO_EXPOSURE_DATA",
                hazard_intensity=60.0,
                hazard_intensity_unit="m/s",
            )
        )
        rows.append(
            _row(
                fid,
                BUSAN,
                "HW",
                K,
                pt,
                lat,
                lon,
                calculation_status="HAZARD_ONLY",
                hazard_intensity=34.06,
                hazard_intensity_unit="degC",
            )
        )
    rows.append(
        _row(
            SEOUL,
            SEOUL,
            "RF",
            C,
            rp,
            s.latitude,
            s.longitude,
            calculation_status="NO_EXPOSURE_DATA",
            hazard_intensity=6.972,
            hazard_intensity_unit="m",
        )
    )
    return rows


def test_comparison_frame_shows_both_anchors_and_not_available() -> None:
    frame = rt.official_office_comparison_frame(_two_anchor_rows())
    assert [line["municipality"] for line in frame] == ["부산광역시", "서울특별시"]
    busan, seoul = frame
    o, b = oo.offices_by_id()[BUSAN], mu.by_id([BUSAN])[0]
    assert (busan["office_lat"], busan["office_lon"]) == (o.latitude, o.longitude)
    assert (busan["representative_lat"], busan["representative_lon"]) == (b.latitude, b.longitude)
    assert (
        busan["coordinate_method"] == "GOVERNMENT_PUBLISHED" and busan["polygon_check"] == "INSIDE"
    )
    assert busan["flood_representative"] == "No asset value (0 m)"
    assert busan["flood_official_office"] == "No asset value (1.5 m)"
    assert busan["heatwave_official_office"] == "Hazard only (34.06 degC)"
    assert busan["distance_between_points_km"] == pytest.approx(
        oo.haversine_km(b.latitude, b.longitude, o.latitude, o.longitude), abs=5e-4
    )
    assert seoul["flood_official_office"] == oo.NOT_AVAILABLE_LABEL
    assert seoul["office_lat"] is None and seoul["distance_between_points_km"] is None
    assert seoul["office_status"] == oo.STATUS_NOT_AVAILABLE
    for line in frame:  # descriptive only — no ranking words anywhere
        text = " ".join(str(v) for v in line.values()).lower()
        assert not any(w in text for w in ("more accurate", "better", "worse", "preferred"))
    assert list(busan)[:17] == [
        "municipality", "office_name", "office_address", "office_lat", "office_lon",
        "coordinate_method", "coordinate_source", "coordinate_source_date", "polygon_check",
        "representative_lat", "representative_lon", "flood_representative",
        "flood_official_office", "tc_representative", "tc_official_office",
        "heatwave_representative", "heatwave_official_office",
    ]  # fmt: skip


def test_two_anchor_frames_label_the_coordinate_type_and_add_the_sheet() -> None:
    frames = rt.export_frames(_two_anchor_rows(), target="MUNICIPALITY")
    assert "official_office_comparison" in frames
    matrix = frames["municipality_risk_matrix"]
    assert [(r["Municipality ID"], r["Coordinate Type"]) for r in matrix] == [
        (BUSAN, "Representative Point"),
        (BUSAN, "Official Office"),
        (SEOUL, "Representative Point"),
    ]
    output = {
        "status": "ok",
        "assessment_target": "MUNICIPALITY",
        "rows": [r.to_dict() for r in _two_anchor_rows()],
        "official_office_dataset": oo.dataset_summary(),
    }
    wb = load_workbook(io.BytesIO(xl.workbook_bytes(output)))
    assert "Official Office Comparison" in wb.sheetnames
    assert wb.sheetnames.index("Official Office Comparison") < wb.sheetnames.index("Methodology")
    info = {r[0]: r[1] for r in wb["Run Info"].iter_rows(values_only=True)}
    assert info["n_municipalities"] == 2  # two anchors of one municipality count once
    assert info["official_office_dataset.official_coordinates_found"] == 1
    assert xl.report_filename(output, "2026-09-27").endswith("_2_Municipalities_20260927.xlsx")


def test_representative_only_frames_are_unchanged() -> None:
    rows = [r for r in _two_anchor_rows() if r.point_type != oo.POINT_TYPE_OFFICIAL_OFFICE]
    frames = rt.export_frames(rows, target="MUNICIPALITY")
    assert "official_office_comparison" not in frames
    assert all("Coordinate Type" not in r for r in frames["municipality_risk_matrix"])
    assert all("Coordinate Type" not in r for r in frames["municipality_summary"])


# --------------------------------------------------------------------------- #
# Build script and API                                                         #
# --------------------------------------------------------------------------- #
def _build_module():  # type: ignore[no-untyped-def]
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "build_official_offices", REPO / "scripts" / "build_official_offices.py"
    )
    assert spec and spec.loader
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


def test_build_takes_only_the_exactly_named_row_of_the_checksummed_file(tmp_path: Path) -> None:
    mod = _build_module()
    body = (
        "행정공공기관명,도로명주소,위도,경도\n"
        "부산광역시청,addr,1.0,2.0\n"
        "부산광역시청 별관,addr2,3.0,4.0\n"
    )
    f = tmp_path / "src.csv"
    f.write_text(body, encoding="utf-8")
    source = {
        "file": "src.csv",
        "encoding": "utf-8",
        "sha256": hashlib.sha256(body.encode()).hexdigest(),
        "name_column": "행정공공기관명",
        "row_name": "부산광역시청",
    }
    assert mod._published_row(source, tmp_path)["위도"] == "1.0"
    with pytest.raises(SystemExit, match="sha256"):
        mod._published_row({**source, "sha256": "0" * 64}, tmp_path)
    with pytest.raises(SystemExit, match="exactly one row"):
        mod._published_row({**source, "row_name": "부산시청"}, tmp_path)


def test_api_serves_offices_and_refuses_office_only_for_unsupported(client) -> None:  # type: ignore[no-untyped-def]
    body = client.get("/api/libraries/municipalities").json()
    assert [o["municipality_id"] for o in body["official_offices"]] == list(oo.TARGETS)
    summary = body["official_office_summary"]
    assert (summary["targets"], summary["official_coordinates_found"]) == (7, 1)
    assert summary["supported_ids"] == [BUSAN]
    assert summary["not_available_label"] == "Not available in V0.3 POC"
    sid = client.post("/api/session").json()["id"]
    url = f"/api/session/{sid}/physical-risk-models"
    base = {"assessment_target": "MUNICIPALITY", "municipality_ids": [SEOUL]}
    r = client.post(url, json={**base, "anchors": ["OFFICIAL_OFFICE_POINT"]})
    assert r.status_code == 400 and "Not available in V0.3 POC" in r.json()["detail"]
    r = client.post(url, json={**base, "anchors": ["CITY_CENTRE"]})
    assert r.status_code == 400 and "anchors" in r.json()["detail"]


def test_anchor_rules_allow_both_anchors_with_unsupported_cities() -> None:
    pytest.importorskip("fastapi")
    from fastapi import HTTPException

    from climaterisk.api.routers.run import _validated_anchors

    assert _validated_anchors(None, [SEOUL]) == ["REPRESENTATIVE_POINT"]
    assert _validated_anchors(list(oo.ANCHORS), [SEOUL, BUSAN]) == list(oo.ANCHORS)
    assert _validated_anchors(["OFFICIAL_OFFICE_POINT"], [BUSAN]) == ["OFFICIAL_OFFICE_POINT"]
    with pytest.raises(HTTPException):
        _validated_anchors(["OFFICIAL_OFFICE_POINT"], [BUSAN, GWANGJU])
