"""Phase 6 — non-expert portfolio views: display copy, recommended models, summary, matrix,
first Excel sheets, CLI words and asset selection.

CLIMADA-free. The rules under test: words never change numbers (the band copy quotes the
configured thresholds); "recommended" never picks a cell that cannot run; the portfolio
summary keeps an unpriced hazard as an empty cell with a status, and a computed zero as 0;
no overall risk level is invented; the Excel first sheets use plain titles and currency /
percent formats; the CLI understands ``flood,typhoon,heatwave`` and ``--models recommended``.
"""

from __future__ import annotations

import importlib.util
import io
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

REPO = Path(__file__).resolve().parents[1]
if str(REPO / "src") not in sys.path:
    sys.path.insert(0, str(REPO / "src"))

from climaterisk.engines.base import PhysicalRiskModelsRequest  # noqa: E402
from climaterisk.physical_risk import display_copy as dc  # noqa: E402
from climaterisk.physical_risk import excel_export as xl  # noqa: E402
from climaterisk.physical_risk import results_table as rt  # noqa: E402
from climaterisk.physical_risk.metrics import (  # noqa: E402
    CalculationStatus,
    ModelId,
    ResultRow,
    load_config,
)

G, C, K = ModelId.GLOBAL_BASELINE.value, ModelId.DATA_API_COUNTRY.value, ModelId.KOREA_LOCAL.value


def _row(fid: str, haz: str, model: str, **kw) -> ResultRow:  # type: ignore[no-untyped-def]
    base = {
        "facility_id": fid,
        "facility_name": f"Facility {fid}",
        "hazard_type": haz,
        "model_id": model,
        "asset_value_usd": 100_000_000.0,
        "requested_scenario": "rcp60",
        "served_scenario": "rcp60",
        "hazard_source": f"synthetic {model}",
    }
    base.update(kw)
    return ResultRow(**base).finalise()  # type: ignore[arg-type]


def _full(fid: str, haz: str, model: str, eal: float, **kw) -> ResultRow:  # type: ignore[no-untyped-def]
    return _row(
        fid,
        haz,
        model,
        eal_usd=eal,
        potential_loss_usd=eal * 96,
        return_period_years=100.0,
        impact_function_id=22 if haz == "RF" else 9,
        hazard_intensity=1.0,
        calculation_status=CalculationStatus.FULL.value,
        **kw,
    )


def _portfolio() -> list[ResultRow]:
    """Three facilities: F1 low/low, F2 medium flood, F3 flood zero (real 0) — heat hazard-only."""
    rows: list[ResultRow] = []
    for fid, flood_eal in (("F1", 75_000.0), ("F2", 200_000.0), ("F3", 0.0)):
        rows += [
            _full(fid, "RF", G, flood_eal),
            _full(fid, "RF", C, flood_eal),
            _row(fid, "RF", K, calculation_status="NOT_IMPLEMENTED"),
            _full(fid, "TC", G, 3_576.0),
            _full(fid, "TC", C, 3_576.0),
            _row(fid, "TC", K, calculation_status="NOT_IMPLEMENTED"),
            _row(fid, "HW", G, calculation_status="NO_HAZARD_DATA"),
            _row(fid, "HW", C, calculation_status="NO_HAZARD_DATA"),
            _row(fid, "HW", K, hazard_intensity=37.1, calculation_status="HAZARD_ONLY"),
        ]
    return rows


# --------------------------------------------------------------------------- #
# Display copy                                                                 #
# --------------------------------------------------------------------------- #
def test_display_copy_quotes_the_configured_thresholds() -> None:
    bands = {b["level"]: b for b in load_config()["risk_levels"]["bands"]}
    copy = dc.risk_level_copy()
    assert (
        copy["Low"]
        == f"Annual expected loss is below {bands['Low']['max_pct']:.2f}% of asset value."
    )
    assert "between 0.10% and 0.50%" in copy["Medium"]
    assert copy["High"].startswith("Annual expected loss is at least 0.50%")
    # every status has a sentence and a short label; the heat sentence names the reason
    for s in CalculationStatus:
        assert s.value in dc.STATUS_COPY and s.value in dc.STATUS_SHORT
    assert "no applicable CLIMADA Impact Function" in dc.STATUS_COPY["HAZARD_ONLY"]
    assert "not domestic source data" in dc.SCOPE_COPY[C]
    assert any("No overall portfolio risk score" in line for line in dc.LIMITATIONS)


def test_recommended_models_follow_readiness_and_never_pick_an_unrunnable_cell() -> None:
    rec = dc.recommended_models()
    assert rec == {"RF": C, "TC": C, "HEAT": K}
    # a registry where local flood is READY flips the recommendation without code changes
    table = {
        "RF": {G: {"status": "READY"}, C: {"status": "READY"}, K: {"status": "READY"}},
        "TC": {
            G: {"status": "READY"},
            C: {"status": "NO_HAZARD_DATA"},
            K: {"status": "NOT_IMPLEMENTED"},
        },
        "HEAT": {
            G: {"status": "NO_HAZARD_DATA"},
            C: {"status": "NO_HAZARD_DATA"},
            K: {"status": "NOT_IMPLEMENTED"},
        },
    }
    rec2 = dc.recommended_models(readiness=table)
    assert rec2 == {"RF": K, "TC": G}  # HEAT omitted: nothing can run
    avail = dc.scope_availability()
    assert avail[K] == {"RF": "unavailable", "TC": "unavailable", "HEAT": "hazard only"}
    assert avail[C]["RF"] == "available" and avail[G]["HEAT"] == "unavailable"
    bundle = dc.display_bundle()
    assert (
        bundle["recommended_models"] == rec
        and "near-zero difference is expected" in bundle["comparison_note"]
    )


# --------------------------------------------------------------------------- #
# Primary rows, summary, matrix, counts                                        #
# --------------------------------------------------------------------------- #
def test_primary_row_prefers_the_recommended_model_but_never_hides_a_priced_one() -> None:
    rows = _portfolio()
    prim = rt.primary_rows(rows, dc.recommended_models())
    assert prim[("F1", "RF")].model_id == C and prim[("F1", "HW")].model_id == K
    # recommended model absent from the run -> fall back to the model that priced it
    only_global = [r for r in rows if r.model_id == G]
    prim_g = rt.primary_rows(only_global, dc.recommended_models())
    assert prim_g[("F1", "RF")].model_id == G
    # an unavailable local row never wins over a priced country row
    assert (
        rt.primary_row(
            [r for r in rows if r.facility_id == "F1" and r.hazard_type == "RF"], K
        ).model_id
        == C
    )


def test_portfolio_summary_keeps_status_empty_cells_and_real_zeros_apart() -> None:
    frame = {
        line["Facility ID"]: line
        for line in rt.portfolio_summary_frame(_portfolio(), dc.recommended_models())
    }
    f1, f2, f3 = frame["F1"], frame["F2"], frame["F3"]
    assert (
        f1["Flood Risk"] == "Low"
        and f1["Flood EAL"] == 75_000.0
        and f1["Flood Data Source"] == "Country"
    )
    assert f2["Flood Risk"] == "Medium" and f2["Flood EAL / Assets"] == pytest.approx(0.2)
    assert f3["Flood Risk"] == "Low" and f3["Flood EAL"] == 0.0  # a computed zero stays 0
    assert f1["Heatwave Status"] == "Hazard only" and f1["Heatwave Tmax p95 (°C)"] == 37.1
    assert "Heatwave Risk" not in f1 and "Heatwave EAL" not in f1  # no financial columns for heat
    assert f1["Overall EAL"] == pytest.approx(75_000.0 + 3_576.0)
    assert f1["Overall EAL / Assets"] == pytest.approx(0.078576)
    assert "Overall Risk" not in f1  # no invented overall level
    # a facility with nothing priced: empty, not zero
    unpriced = rt.portfolio_summary_frame(
        [
            _row("X", "RF", K, calculation_status="NOT_IMPLEMENTED"),
            _row("X", "HW", K, calculation_status="HAZARD_ONLY"),
        ]
    )[0]
    assert unpriced["Flood EAL"] is None and unpriced["Flood Status"] == "Unavailable"
    assert unpriced["Overall EAL"] is None and unpriced["Typhoon Status"] == "Not assessed"


def test_asset_matrix_and_counts_use_levels_or_status_labels_only() -> None:
    rows = _portfolio()
    matrix = {
        m["Facility ID"]: m for m in rt.asset_risk_matrix_frame(rows, dc.recommended_models())
    }
    assert matrix["F2"] == {
        "Facility ID": "F2",
        "Facility": "Facility F2",
        "Flood": "Medium",
        "Typhoon": "Low",
        "Heatwave": "Hazard only",
        "Overall": "2 of 3 calculated",
    }
    counts = rt.summary_counts(rows, dc.recommended_models())
    assert counts == {
        "assets_analyzed": 3,
        "high_risk": 0,
        "medium_risk": 1,
        "low_risk": 5,
        "hazard_only": 3,
        "not_available": 0,
        "errors": 0,
    }
    # statuses that are "not available" are counted, not hidden
    c2 = rt.summary_counts([_row("X", "RF", C, calculation_status="SCENARIO_MISMATCH")])
    assert c2["not_available"] == 1 and c2["assets_analyzed"] == 1


# --------------------------------------------------------------------------- #
# Excel first sheets                                                           #
# --------------------------------------------------------------------------- #
def _output() -> dict:  # type: ignore[type-arg]
    from climaterisk.physical_risk.models import readiness

    return {
        "status": "ok",
        "climate_scenario": "rcp60",
        "target_year": 2050,
        "country": "KOR",
        "rows": [r.to_dict() for r in _portfolio()],
        "readiness": readiness(),
        "impact_function_fixed": {"RF": True, "TC": True, "HW": True},
    }


def test_excel_first_sheets_are_plain_and_formatted() -> None:
    wb = load_workbook(io.BytesIO(xl.workbook_bytes(_output())))
    assert wb.sheetnames[:3] == ["Portfolio Summary", "Asset Risk Matrix", "Hazard Results"]
    ws = wb["Portfolio Summary"]
    header = [c.value for c in ws[1]]
    for wanted in (
        "Facility",
        "Asset Value",
        "Flood Risk",
        "Flood EAL",
        "Flood EAL / Assets",
        "Typhoon Risk",
        "Heatwave Status",
        "Overall EAL",
        "Overall EAL / Assets",
    ):
        assert wanted in header, wanted
    for technical in (
        "haz_type",
        "impf_id",
        "impact_function_id",
        "MDD",
        "PAA",
        "hazard_adapter",
        "model_id",
    ):
        assert technical not in header, technical
    col = {h: i + 1 for i, h in enumerate(header)}
    assert ws.cell(row=2, column=col["Asset Value"]).number_format == '"$"#,##0'
    assert ws.cell(row=2, column=col["Flood EAL"]).number_format == '"$"#,##0'
    assert ws.cell(row=2, column=col["Flood EAL / Assets"]).number_format == '0.000"%"'
    # F3's flood EAL is a real zero -> 0 in the cell; heat money cells do not exist at all
    rows = {ws.cell(row=i, column=col["Facility ID"]).value: i for i in range(2, ws.max_row + 1)}
    assert ws.cell(row=rows["F3"], column=col["Flood EAL"]).value == 0
    assert ws.cell(row=rows["F1"], column=col["Heatwave Status"]).value == "Hazard only"
    # risk columns carry conditional formatting (a visual aid, not a rule)
    assert any("Flood Risk" in header and r for r in ws.conditional_formatting)
    matrix = wb["Asset Risk Matrix"]
    assert [c.value for c in matrix[1]] == [
        "Facility ID",
        "Facility",
        "Flood",
        "Typhoon",
        "Heatwave",
        "Overall",
    ]
    assert matrix.freeze_panes == "A2"
    meth = {r[0]: r[1] for r in wb["Methodology"].iter_rows(min_row=2, values_only=True)}
    assert meth["Risk Level — Low"].startswith("Annual expected loss is below 0.10%")
    assert "no applicable CLIMADA Impact Function" in meth["Calculation status — HAZARD_ONLY"]
    assert any(k.startswith("Limitations — ") for k in meth)
    for section in (
        "What was assessed",
        "Hazards — Flood",
        "CLIMADA method",
        "Impact Functions",
        "EAL",
        "Potential Loss",
        "Climate Change Multiplier",
    ):
        assert section in meth, section
    # first-sheet aggregates the brief asks for
    assert {"Priced Hazard Count", "Hazard-only Count", "Not Available Count"} <= set(header)
    assert ws.cell(row=rows["F1"], column=col["Priced Hazard Count"]).value == 2
    assert ws.cell(row=rows["F1"], column=col["Hazard-only Count"]).value == 1


# --------------------------------------------------------------------------- #
# CLI words and asset selection                                                #
# --------------------------------------------------------------------------- #
def _cli():  # type: ignore[no-untyped-def]
    spec = importlib.util.spec_from_file_location(
        "physical_risk_batch", REPO / "scripts" / "physical_risk_batch.py"
    )
    assert spec is not None and spec.loader is not None
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_cli_understands_plain_words_and_recommended_reads_readiness() -> None:
    cli = _cli()
    assert cli.parse_hazards(["flood,typhoon,heatwave"]) == ["RF", "TC", "HEAT"]
    assert cli.parse_hazards(["RF", "heat"]) == ["RF", "HEAT"]
    with pytest.raises(ValueError, match="unknown hazard"):
        cli.parse_hazards(["earthquake"])
    assert cli.parse_models(["recommended"], ["RF", "TC", "HEAT"]) == [C, K]
    assert cli.parse_models(["recommended"], ["RF"]) == [C]
    assert cli.parse_models(["all"], ["RF"]) == [G, C, K]
    assert cli.parse_models(["global,country"], ["RF"]) == [G, C]
    with pytest.raises(ValueError, match="unknown model"):
        cli.parse_models(["regional"], ["RF"])


def test_cli_csv_validation_speaks_plainly(tmp_path: Path) -> None:
    cli = _cli()
    dup = tmp_path / "dup.csv"
    dup.write_text(
        "facility_id,lat,lon,asset_value_usd\nA,37.5,127,1\nA,35.1,129,1\n", encoding="utf-8"
    )
    with pytest.raises(ValueError, match="duplicate facility_id"):
        cli.read_facilities(dup)
    rng = tmp_path / "rng.csv"
    rng.write_text("facility_id,lat,lon,asset_value_usd\nA,137.5,127,1\n", encoding="utf-8")
    with pytest.raises(ValueError, match="out of range"):
        cli.read_facilities(rng)
    neg = tmp_path / "neg.csv"
    neg.write_text("facility_id,lat,lon,asset_value_usd\nA,37.5,127,-5\n", encoding="utf-8")
    with pytest.raises(ValueError, match="negative"):
        cli.read_facilities(neg)
    ok = tmp_path / "ok.csv"
    ok.write_text(
        "facility_id,facility_name,lat,lon,asset_value_usd\nA,Alpha,37.5,127,1000\n",
        encoding="utf-8",
    )
    assert cli.read_facilities(ok)[0]["facility_name"] == "Alpha"


def test_request_selects_a_subset_of_assets_and_refuses_unknown_ids(client) -> None:  # type: ignore[no-untyped-def]
    from climaterisk.core.entities import Asset, Portfolio

    pf = Portfolio(
        assets=[
            Asset(id="a1", name="One", lat=37.5, lon=127.0, value=1e8),
            Asset(id="a2", name="Two", lat=35.1, lon=129.0, value=6e7),
            Asset(id="a3", name="Three", lat=37.4, lon=126.7, value=8e7),
        ]
    )
    req = PhysicalRiskModelsRequest.from_portfolio(pf, facility_ids=["a1", "a3"])
    assert [f.facility_id for f in req.facilities] == ["a1", "a3"]
    assert len(PhysicalRiskModelsRequest.from_portfolio(pf).facilities) == 3
    assert PhysicalRiskModelsRequest.from_portfolio(pf, facility_ids=[]).facilities == []
    # the route validates ids against the session before spawning anything
    r = client.post("/api/session/nope/physical-risk-models", json={"facility_ids": ["a1"]})
    assert r.status_code == 404
    r = client.get("/api/libraries/physical-risk-models")
    assert r.status_code == 200 and r.json()["display"]["recommended_models"]["HEAT"] == K


def test_run_history_lists_physical_risk_runs_newest_first(client, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    """Run history: the store filters by kind and the route 404s an unknown session."""
    from climaterisk.runs.store import RunStore

    assert client.get("/api/session/nope/runs").status_code == 404
    store = RunStore(tmp_path / "runs.db")
    older = store.create("r1", "s1", "rcp85", ["physical_risk_models"])
    store.create("r2", "s1", "rcp85", ["tropical_cyclone"])  # a legacy physical run — not listed
    newer = store.create("r3", "s1", "rcp85", ["physical_risk_models"])
    listed = store.list_by_kind("s1", "physical_risk_models", limit=10)
    assert [r.id for r in listed] in (
        [newer.id, older.id],
        [older.id, newer.id],
    )  # same-second timestamps
    assert all(r.perils == ["physical_risk_models"] for r in listed)
    assert store.list_by_kind("s1", "physical_risk_models", limit=1)[0].id in (newer.id, older.id)
    assert store.list_by_kind("s2", "physical_risk_models") == []


def test_report_filename_uses_asset_count_and_run_date() -> None:
    """Physical_Risk_Report_<N>_Assets_<YYYYMMDD>.xlsx — derived from the run, not typed."""
    rows = [{"facility_id": f} for f in ("a", "b", "a", "c")]
    n_assets = len({r["facility_id"] for r in rows})
    stamp = "2026-09-24T10:00:00Z"[:10].replace("-", "")
    assert (
        f"Physical_Risk_Report_{n_assets}_Assets_{stamp}.xlsx"
        == "Physical_Risk_Report_3_Assets_20260924.xlsx"
    )


def test_display_bundle_carries_the_final_copy() -> None:
    b = dc.display_bundle()
    assert "not an overall portfolio risk score" in b["summary_note"]
    assert "pipeline-consistency" in b["comparison_purpose"]
    assert any("Korea asset-value exposure" in line for line in b["limitations"])
    assert all(
        "modeled" in dc.HAZARD_COPY[k].lower() or "Modeled" in dc.HAZARD_COPY[k]
        for k in ("RF", "TC")
    )
    for banned in ("safe", "no risk", "guaranteed", "actual flood depth", "actual expected damage"):
        assert banned not in " ".join(b["limitations"]).lower()
        assert banned not in " ".join(dc.STATUS_COPY.values()).lower()


def test_hazard_results_sheet_carries_the_plain_data_source_column() -> None:
    rows = _portfolio()
    frame = rt.hazard_results_frame(rows)
    assert "data_source" in frame[0] and list(frame[0]).index("data_source") == list(
        rt.HAZARD_RESULTS_COLUMNS
    ).index("data_source")
    assert {r["data_source"] for r in frame} == {"Global", "Country", "Local"}
    for c in (
        "facility_id",
        "facility_name",
        "hazard_type",
        "model_id",
        "data_source",
        "risk_level",
        "risk_level_criteria",
        "probability",
        "potential_loss_usd",
        "eal_usd",
        "eal_as_pct_of_assets",
        "return_period_years",
        "climate_change_multiplier",
        "calculation_status",
    ):
        assert c in rt.HAZARD_RESULTS_COLUMNS, c
